"""P1b-2: connector documents go through the same (A1) extractors as uploads.

Live finding (2026-10-05, MinIO source with mixed formats): a .pptx and a .zip
object were decoded as UTF-8 text and failed in Postgres with "invalid byte
sequence for encoding UTF8: 0x00" — the connector pipeline had no PPTX or ZIP
extractor, used the old regex HTML stripper (page chrome indexed, title twice)
and the older DOCX / PDF parsers (no tracked insertions / headers, no per-page
OCR), while file upload had all of them.
"""

from __future__ import annotations

import io
import zipfile
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

from app.ingestion.content_classifier import ContentType
from app.ingestion.parser_registry import DocumentParseError, ParserRegistry
from app.ingestion.pipeline import IngestionPipeline
from app.ingestion.source_config import RawDocument, SourceConfig, SourceFamily

PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _pptx(title: str, body: str, notes: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = title
    s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2)).text_frame.text = body
    s.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def _docx_with_header(body: str, header: str) -> bytes:
    from docx import Document

    d = Document()
    d.sections[0].header.paragraphs[0].text = header
    d.add_paragraph(body)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _zip(members: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return buf.getvalue()


async def _parse(data: bytes, name: str, mime: str = "", ct: ContentType = ContentType.TEXT
                 ) -> tuple[str, dict[str, object]]:
    return await ParserRegistry().parse_bytes_async(data, ct, filename=name, mime_type=mime)


@pytest.mark.asyncio
async def test_pptx_by_name_or_mime_reads_slides_and_notes() -> None:
    data = _pptx("Kestrel Loop", "Ranchi target 96.2 percent", "Speaker notes here")
    for name, mime in (("deck.pptx", ""), ("deck", PPTX)):
        text, _ = await _parse(data, name, mime)
        assert "Ranchi target 96.2 percent" in text
        assert "Speaker notes here" in text
        assert "\x00" not in text


@pytest.mark.asyncio
async def test_zip_members_are_extracted_with_their_paths() -> None:
    data = _zip({
        "handover/vapi.md": b"# Vapi\n\nThe Vapi bay key is held by Ishaan Rao.\n",
        "deck.pptx": _pptx("Deck", "Slide body inside the archive", "n"),
        "__MACOSX/._junk": b"\x00\x01",
        "legacy.doc": b"\xd0\xcf\x11\xe0" + b"\x00" * 64,
    })
    text, meta = await _parse(data, "bundle.zip", "application/zip")
    assert "handover/vapi.md" in text and "Ishaan Rao" in text
    assert "Slide body inside the archive" in text
    assert "\x00" not in text
    assert any("legacy.doc" in str(s) for s in meta.get("archive_members_skipped", []))  # type: ignore[union-attr]


@pytest.mark.asyncio
async def test_zip_bomb_is_refused() -> None:
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("big.txt", b"0" * (60 * 1024 * 1024))
    with pytest.raises(DocumentParseError):
        await _parse(bomb.getvalue(), "bomb.zip", "application/zip")


@pytest.mark.asyncio
async def test_unknown_binary_is_a_parse_failure_not_text() -> None:
    with pytest.raises(DocumentParseError, match="binary"):
        await _parse(b"\x7fELF\x02\x01\x01" + b"\x00" * 200 + b"payload", "tool.bin",
                     "binary/octet-stream")


@pytest.mark.asyncio
async def test_html_drops_page_chrome() -> None:
    html = (b"<html><head><title>T</title></head><body><nav>Home | About | Login</nav>"
            b"<main><h1>Bulletin</h1><p>Gate pass scheme Orchid-9 starts in week 44.</p></main>"
            b"<footer>Subscribe to our newsletter</footer></body></html>")
    text, _ = await _parse(html, "b.html", "text/html", ContentType.HTML)
    assert "Orchid-9" in text
    assert "newsletter" not in text and "Login" not in text


@pytest.mark.asyncio
async def test_docx_headers_are_read_like_uploads() -> None:
    data = _docx_with_header("Body clause about reefer probes.", "CONFIDENTIAL Halvorsen")
    text, _ = await _parse(data, "c.docx", ct=ContentType.DOCX)
    assert "Body clause about reefer probes." in text
    assert "CONFIDENTIAL Halvorsen" in text


@pytest.mark.asyncio
async def test_corrupt_pdf_fails_with_a_reason() -> None:
    with pytest.raises(DocumentParseError):
        await _parse(b"%PDF-1.4\n1 0 obj << /Type /Catalog >>\n% truncated", "broken.pdf",
                     "application/pdf", ContentType.PDF)


@pytest.mark.asyncio
async def test_pipeline_indexes_a_pptx_object_from_a_connector() -> None:
    kb = MagicMock()
    kb.exists_by_hash = AsyncMock(return_value=False)
    kb.ingest_chunks_async = AsyncMock(return_value=["c1"])
    pipeline = IngestionPipeline(knowledge_store=kb, embedder=MagicMock())
    raw = RawDocument(
        doc_id="s3://b/decks/kestrel.pptx", source_id="s1", tenant_id="t1",
        content=_pptx("Kestrel Loop", "Project Kestrel Loop sets the Ranchi depot target "
                      "at 96.2 percent on-time dispatch for the quarter.", "notes"),
        content_type=PPTX, title="kestrel.pptx", source_url="s3://b/decks/kestrel.pptx",
    )
    cfg = SourceConfig(source_id="s1", tenant_id="t1", name="s",
                       family=SourceFamily.OBJECT_STORAGE, source_type="minio",
                       collection_id="c", pii_action="allow", min_quality_score=0.0)
    with patch("app.providers.base.embed_texts",
               AsyncMock(side_effect=lambda texts, **_: [[0.1] * 8 for _ in texts])):
        result = await pipeline.ingest(raw, cfg)
    assert result.status == "indexed", result.error
    chunks = kb.ingest_chunks_async.call_args[0][0]
    assert any("96.2 percent" in c.content for c in chunks)
    assert all("\x00" not in c.content for c in chunks)
