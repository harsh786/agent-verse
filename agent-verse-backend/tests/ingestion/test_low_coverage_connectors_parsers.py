"""Coverage tests for low-coverage ingestion connectors and parsers.

Targets (per coverage audit):
  - app/ingestion/connectors/gcs_connector.py
  - app/ingestion/connectors/rss_connector.py
  - app/ingestion/connectors/github_connector.py
  - app/ingestion/connectors/azure_blob_connector.py
  - app/ingestion/connectors/slack_connector.py
  - app/ingestion/connectors/pdf_file_connector.py (PDFFileConnector + DOCXFileConnector)
  - app/ingestion/parsers/parquet_parser.py
  - app/ingestion/parsers/avro_parser.py
  - app/ingestion/parsers/excel_parser.py

All connector tests mock the external SDK / HTTP layer — no live network, DB, or
cloud credentials required. Parser tests exercise the real libraries (pyarrow,
fastavro, openpyxl, feedparser — installed in this environment) for the
library-dependent code paths, plus the ImportError fallback paths via
sys.modules patching.
"""
from __future__ import annotations

import io
from datetime import datetime, UTC
from unittest.mock import AsyncMock, MagicMock, patch


from app.ingestion.source_config import SourceConfig, SourceFamily


def _config(source_type: str, conn_config: dict | None = None) -> SourceConfig:
    return SourceConfig(
        source_id="src-1",
        tenant_id="tenant-1",
        name="Test Source",
        family=SourceFamily.OBJECT_STORAGE,
        source_type=source_type,
        connection_config=conn_config or {},
    )


def _fake_module_tree(dotted: str, leaf: object | None = None) -> dict[str, object]:
    """Build sys.modules entries so `from a.b import c` resolves to a mock.

    Returns a dict suitable for `patch.dict("sys.modules", ...)` that makes the
    dotted module path importable, with the leaf module set to `leaf` (or a
    fresh MagicMock), and each parent package exposing the right attribute.
    """
    parts = dotted.split(".")
    leaf_obj = leaf if leaf is not None else MagicMock()
    modules: dict[str, object] = {dotted: leaf_obj}
    child = leaf_obj
    for i in range(len(parts) - 1, 0, -1):
        parent_name = ".".join(parts[:i])
        parent_mod = MagicMock()
        setattr(parent_mod, parts[i], child)
        modules[parent_name] = parent_mod
        child = parent_mod
    return modules


async def _collect(agen) -> list:
    return [item async for item in agen]


# ═══════════════════════════════════════════════════════════════════════════
# GCSConnector
# ═══════════════════════════════════════════════════════════════════════════

class TestGCSConnectorValidateConnection:
    def test_bucket_exists_ok(self):
        from app.ingestion.connectors.gcs_connector import GCSConnector

        storage_mock = MagicMock()
        client = MagicMock()
        storage_mock.Client.return_value = client
        bucket = MagicMock()
        client.bucket.return_value = bucket
        bucket.exists.return_value = True
        client.list_blobs.return_value = [MagicMock()]

        config = _config("gcs", {"bucket": "my-bucket"})
        with patch.dict("sys.modules", _fake_module_tree("google.cloud.storage", storage_mock)):
            result = _run(GCSConnector().validate_connection(config))

        assert result.ok is True
        assert result.metadata["bucket"] == "my-bucket"
        assert result.metadata["sample_objects"] == 1

    def test_bucket_missing_not_ok(self):
        from app.ingestion.connectors.gcs_connector import GCSConnector

        storage_mock = MagicMock()
        client = MagicMock()
        storage_mock.Client.return_value = client
        bucket = MagicMock()
        client.bucket.return_value = bucket
        bucket.exists.return_value = False

        config = _config("gcs", {"bucket": "missing-bucket"})
        with patch.dict("sys.modules", _fake_module_tree("google.cloud.storage", storage_mock)):
            result = _run(GCSConnector().validate_connection(config))

        assert result.ok is False
        assert "not found" in result.error

    def test_service_account_json_dict_uses_tempfile(self):
        from app.ingestion.connectors.gcs_connector import GCSConnector

        storage_mock = MagicMock()
        client = MagicMock()
        storage_mock.Client.from_service_account_json.return_value = client
        bucket = MagicMock()
        client.bucket.return_value = bucket
        bucket.exists.return_value = True
        client.list_blobs.return_value = []

        config = _config(
            "gcs", {"bucket": "b1", "service_account_json": {"type": "service_account"}}
        )
        with patch.dict("sys.modules", _fake_module_tree("google.cloud.storage", storage_mock)):
            result = _run(GCSConnector().validate_connection(config))

        assert result.ok is True
        storage_mock.Client.from_service_account_json.assert_called_once()

    def test_exception_returns_not_ok(self):
        from app.ingestion.connectors.gcs_connector import GCSConnector

        storage_mock = MagicMock()
        storage_mock.Client.side_effect = RuntimeError("boom")

        config = _config("gcs", {"bucket": "b1"})
        with patch.dict("sys.modules", _fake_module_tree("google.cloud.storage", storage_mock)):
            result = _run(GCSConnector().validate_connection(config))

        assert result.ok is False
        assert "boom" in result.error


class TestGCSConnectorGetDelta:
    def _blob(self, name, updated, content=b"hello", content_type="text/plain", size=5):
        blob = MagicMock()
        blob.name = name
        blob.updated = updated
        blob.content_type = content_type
        blob.size = size
        blob.download_as_bytes.return_value = content
        return blob

    def test_import_error_yields_nothing(self):
        from app.ingestion.connectors.gcs_connector import GCSConnector

        config = _config("gcs", {"bucket": "b1"})
        with patch.dict("sys.modules", {"google.cloud.storage": None}):
            docs = _run(_collect(GCSConnector().get_delta(config, None)))
        assert docs == []

    def test_yields_documents_and_advances_cursor(self):
        from app.ingestion.connectors.gcs_connector import GCSConnector

        storage_mock = MagicMock()
        client = MagicMock()
        storage_mock.Client.return_value = client
        blob1 = self._blob("a.txt", datetime(2026, 1, 1, tzinfo=UTC))
        blob2 = self._blob("b.txt", datetime(2026, 1, 2, tzinfo=UTC))
        client.list_blobs.return_value = [blob1, blob2]

        config = _config("gcs", {"bucket": "b1", "prefix": "p/"})
        with patch.dict("sys.modules", _fake_module_tree("google.cloud.storage", storage_mock)):
            docs = _run(_collect(GCSConnector().get_delta(config, None)))

        assert len(docs) == 2
        doc, cursor = docs[-1]
        assert doc.source_url == "gs://b1/b.txt"
        assert doc.content == b"hello"
        assert cursor == "2026-01-02T00:00:00+00:00"

    def test_skips_blob_older_than_cursor(self):
        from app.ingestion.connectors.gcs_connector import GCSConnector

        storage_mock = MagicMock()
        client = MagicMock()
        storage_mock.Client.return_value = client
        old_blob = self._blob("old.txt", datetime(2020, 1, 1, tzinfo=UTC))
        client.list_blobs.return_value = [old_blob]

        config = _config("gcs", {"bucket": "b1"})
        with patch.dict("sys.modules", _fake_module_tree("google.cloud.storage", storage_mock)):
            docs = _run(_collect(
                GCSConnector().get_delta(config, "2025-01-01T00:00:00+00:00")
            ))
        assert docs == []

    def test_skips_excluded_pattern(self):
        from app.ingestion.connectors.gcs_connector import GCSConnector

        storage_mock = MagicMock()
        client = MagicMock()
        storage_mock.Client.return_value = client
        blob = self._blob("secret.env", datetime(2026, 1, 1, tzinfo=UTC))
        client.list_blobs.return_value = [blob]

        config = _config("gcs", {"bucket": "b1"})
        config.exclude_patterns = ["*.env"]
        with patch.dict("sys.modules", _fake_module_tree("google.cloud.storage", storage_mock)):
            docs = _run(_collect(GCSConnector().get_delta(config, None)))
        assert docs == []

    def test_download_exception_is_skipped_not_raised(self):
        from app.ingestion.connectors.gcs_connector import GCSConnector

        storage_mock = MagicMock()
        client = MagicMock()
        storage_mock.Client.return_value = client
        bad_blob = self._blob("bad.txt", datetime(2026, 1, 1, tzinfo=UTC))
        bad_blob.download_as_bytes.side_effect = RuntimeError("network error")
        good_blob = self._blob("good.txt", datetime(2026, 1, 2, tzinfo=UTC))
        client.list_blobs.return_value = [bad_blob, good_blob]

        config = _config("gcs", {"bucket": "b1"})
        with patch.dict("sys.modules", _fake_module_tree("google.cloud.storage", storage_mock)):
            docs = _run(_collect(GCSConnector().get_delta(config, None)))
        assert len(docs) == 1
        assert docs[0][0].source_url == "gs://b1/good.txt"


# ═══════════════════════════════════════════════════════════════════════════
# RSSConnector
# ═══════════════════════════════════════════════════════════════════════════

class TestRSSConnectorValidateConnection:
    def test_import_error(self):
        from app.ingestion.connectors.rss_connector import RSSConnector

        config = _config("rss", {"url": "http://example.com/feed"})
        with patch.dict("sys.modules", {"feedparser": None}):
            result = _run(RSSConnector().validate_connection(config))
        assert result.ok is False
        assert "feedparser" in result.error

    def test_success(self):
        from app.ingestion.connectors.rss_connector import RSSConnector

        config = _config("rss", {"url": "http://example.com/feed"})
        with patch("feedparser.parse") as mock_parse:
            mock_feed = MagicMock()
            mock_feed.bozo = False
            mock_feed.entries = [MagicMock(), MagicMock()]
            mock_feed.feed = {"title": "My Feed"}
            mock_parse.return_value = mock_feed
            result = _run(RSSConnector().validate_connection(config))
        assert result.ok is True
        assert result.metadata["title"] == "My Feed"
        assert result.metadata["entries"] == 2

    def test_bozo_with_no_entries_raises(self):
        from app.ingestion.connectors.rss_connector import RSSConnector

        config = _config("rss", {"url": "http://bad.example.com/feed"})
        with patch("feedparser.parse") as mock_parse:
            mock_feed = MagicMock()
            mock_feed.bozo = True
            mock_feed.bozo_exception = ValueError("malformed xml")
            mock_feed.entries = []
            mock_parse.return_value = mock_feed
            result = _run(RSSConnector().validate_connection(config))
        assert result.ok is False
        assert "malformed xml" in result.error


class TestRSSConnectorGetDelta:
    def test_import_error_yields_nothing(self):
        from app.ingestion.connectors.rss_connector import RSSConnector

        config = _config("rss", {"url": "http://example.com/feed"})
        with patch.dict("sys.modules", {"feedparser": None}):
            docs = _run(_collect(RSSConnector().get_delta(config, None)))
        assert docs == []

    def test_skips_entries_older_than_cursor(self):
        from app.ingestion.connectors.rss_connector import RSSConnector

        config = _config("rss", {"url": "http://example.com/feed"})
        connector = RSSConnector()

        old_entry = MagicMock()
        old_entry.get = lambda k, d="": {
            "id": "1", "title": "Old", "link": "http://e.com/1",
            "published": "2020-01-01T00:00:00Z", "summary": "old",
        }.get(k, d)
        new_entry = MagicMock()
        new_entry.get = lambda k, d="": {
            "id": "2", "title": "New", "link": "http://e.com/2",
            "published": "2027-01-01T00:00:00Z", "summary": "new",
        }.get(k, d)

        with patch("feedparser.parse") as mock_parse:
            mock_feed = MagicMock()
            mock_feed.entries = [old_entry, new_entry]
            mock_parse.return_value = mock_feed
            docs = _run(_collect(
                connector.get_delta(config, "2025-01-01T00:00:00Z")
            ))

        assert len(docs) == 1
        assert b"New" in docs[0][0].content

    def test_respects_max_entries(self):
        from app.ingestion.connectors.rss_connector import RSSConnector

        config = _config("rss", {"url": "http://example.com/feed", "max_entries": 1})
        connector = RSSConnector()

        def _mk(i):
            e = MagicMock()
            e.get = lambda k, d="", i=i: {
                "id": str(i), "title": f"T{i}", "link": f"http://e.com/{i}",
                "published": f"2026-01-0{i}T00:00:00Z", "summary": "s",
            }.get(k, d)
            return e

        with patch("feedparser.parse") as mock_parse:
            mock_feed = MagicMock()
            mock_feed.entries = [_mk(1), _mk(2), _mk(3)]
            mock_parse.return_value = mock_feed
            docs = _run(_collect(connector.get_delta(config, None)))

        assert len(docs) == 1


# ═══════════════════════════════════════════════════════════════════════════
# GitHubConnector
# ═══════════════════════════════════════════════════════════════════════════

class TestGitHubConnectorValidateConnection:
    def test_success(self):
        from app.ingestion.connectors.github_connector import GitHubConnector

        config = _config("github", {"token": "tok"})
        with patch("httpx.AsyncClient") as mock_cls:
            resp = MagicMock()
            resp.status_code = 200
            resp.json = MagicMock(return_value={"login": "octocat"})
            client = AsyncMock()
            client.get = AsyncMock(return_value=resp)
            client.__aenter__ = AsyncMock(return_value=client)
            client.__aexit__ = AsyncMock(return_value=False)
            mock_cls.return_value = client
            result = _run(GitHubConnector().validate_connection(config))
        assert result.ok is True
        assert result.metadata["login"] == "octocat"

    def test_non_200_not_ok(self):
        from app.ingestion.connectors.github_connector import GitHubConnector

        config = _config("github", {"token": "bad"})
        with patch("httpx.AsyncClient") as mock_cls:
            resp = MagicMock()
            resp.status_code = 401
            client = AsyncMock()
            client.get = AsyncMock(return_value=resp)
            client.__aenter__ = AsyncMock(return_value=client)
            client.__aexit__ = AsyncMock(return_value=False)
            mock_cls.return_value = client
            result = _run(GitHubConnector().validate_connection(config))
        assert result.ok is False
        assert "401" in result.error

    def test_exception_not_ok(self):
        from app.ingestion.connectors.github_connector import GitHubConnector

        config = _config("github", {"token": "tok"})
        with patch("httpx.AsyncClient", side_effect=RuntimeError("conn refused")):
            result = _run(GitHubConnector().validate_connection(config))
        assert result.ok is False
        assert "conn refused" in result.error


class TestGitHubConnectorGetDelta:
    def test_yields_chunks_as_documents(self):
        from app.ingestion.connectors.github_connector import GitHubConnector

        config = _config("github", {"token": "tok", "repos": ["acme/repo"]})
        chunks = [
            {
                "content": "print('hi')",
                "source_url": "https://github.com/acme/repo/blob/main/a.py",
                "source_doc_id": "acme/repo/a.py",
                "metadata": {"path": "a.py"},
            }
        ]
        with patch(
            "app.knowledge.ingestors.github_ingestor.GitHubIngestor"
        ) as MockIngestor:
            MockIngestor.return_value.ingest_repo = AsyncMock(return_value=chunks)
            docs = _run(_collect(GitHubConnector().get_delta(config, None)))

        assert len(docs) == 1
        doc, cursor = docs[0]
        assert doc.content == b"print('hi')"
        assert doc.doc_id == "acme/repo_acme/repo/a.py"

    def test_include_code_false_skips_ingestion(self):
        from app.ingestion.connectors.github_connector import GitHubConnector

        config = _config(
            "github", {"token": "tok", "repos": ["acme/repo"], "include_code": False}
        )
        with patch(
            "app.knowledge.ingestors.github_ingestor.GitHubIngestor"
        ) as MockIngestor:
            MockIngestor.return_value.ingest_repo = AsyncMock(return_value=[])
            docs = _run(_collect(GitHubConnector().get_delta(config, None)))
        assert docs == []
        MockIngestor.return_value.ingest_repo.assert_not_called()

    def test_repo_error_is_caught_and_continues(self):
        from app.ingestion.connectors.github_connector import GitHubConnector

        config = _config(
            "github", {"token": "tok", "repos": ["acme/broken", "acme/ok"]}
        )
        good_chunk = [{"content": "ok", "source_url": "u", "source_doc_id": "id1", "metadata": {}}]
        with patch(
            "app.knowledge.ingestors.github_ingestor.GitHubIngestor"
        ) as MockIngestor:
            MockIngestor.return_value.ingest_repo = AsyncMock(
                side_effect=[RuntimeError("404"), good_chunk]
            )
            docs = _run(_collect(GitHubConnector().get_delta(config, None)))
        assert len(docs) == 1
        assert docs[0][0].content == b"ok"


# ═══════════════════════════════════════════════════════════════════════════
# AzureBlobConnector
# ═══════════════════════════════════════════════════════════════════════════

class TestAzureBlobConnectorValidateConnection:
    def test_connection_string_success(self):
        from app.ingestion.connectors.azure_blob_connector import AzureBlobConnector

        blob_mod = MagicMock()
        client = MagicMock()
        blob_mod.BlobServiceClient.from_connection_string.return_value = client
        cc = MagicMock()
        client.get_container_client.return_value = cc
        cc.get_container_properties.return_value = {"lease": {"state": "available"}}

        config = _config("azure_blob", {"connection_string": "DefaultEndpoints...", "container": "c1"})
        with patch.dict("sys.modules", _fake_module_tree("azure.storage.blob", blob_mod)):
            result = _run(AzureBlobConnector().validate_connection(config))
        assert result.ok is True
        assert result.metadata["container"] == "c1"

    def test_account_key_success(self):
        from app.ingestion.connectors.azure_blob_connector import AzureBlobConnector

        blob_mod = MagicMock()
        client = MagicMock()
        blob_mod.BlobServiceClient.return_value = client
        cc = MagicMock()
        client.get_container_client.return_value = cc
        cc.get_container_properties.return_value = {}

        config = _config(
            "azure_blob",
            {"account_name": "acct", "account_key": "key", "container": "c1"},
        )
        with patch.dict("sys.modules", _fake_module_tree("azure.storage.blob", blob_mod)):
            result = _run(AzureBlobConnector().validate_connection(config))
        assert result.ok is True

    def test_import_error(self):
        from app.ingestion.connectors.azure_blob_connector import AzureBlobConnector

        config = _config("azure_blob", {"container": "c1"})
        with patch.dict("sys.modules", {"azure.storage.blob": None}):
            result = _run(AzureBlobConnector().validate_connection(config))
        assert result.ok is False
        assert "not installed" in result.error

    def test_exception_not_ok(self):
        from app.ingestion.connectors.azure_blob_connector import AzureBlobConnector

        blob_mod = MagicMock()
        blob_mod.BlobServiceClient.from_connection_string.side_effect = RuntimeError("auth failed")

        config = _config("azure_blob", {"connection_string": "x", "container": "c1"})
        with patch.dict("sys.modules", _fake_module_tree("azure.storage.blob", blob_mod)):
            result = _run(AzureBlobConnector().validate_connection(config))
        assert result.ok is False
        assert "auth failed" in result.error


class TestAzureBlobConnectorGetDelta:
    def _blob(self, name, last_modified, content_type=None, size=5):
        blob = MagicMock()
        blob.name = name
        blob.last_modified = last_modified
        blob.size = size
        if content_type is None:
            blob.content_settings = None
        else:
            blob.content_settings = MagicMock(content_type=content_type)
        return blob

    def test_import_error_yields_nothing(self):
        from app.ingestion.connectors.azure_blob_connector import AzureBlobConnector

        config = _config("azure_blob", {"container": "c1"})
        with patch.dict("sys.modules", {"azure.storage.blob": None}):
            docs = _run(_collect(AzureBlobConnector().get_delta(config, None)))
        assert docs == []

    def test_yields_documents_with_content_type_fallback(self):
        from app.ingestion.connectors.azure_blob_connector import AzureBlobConnector

        blob_mod = MagicMock()
        service = MagicMock()
        blob_mod.BlobServiceClient.from_connection_string.return_value = service
        cc = MagicMock()
        service.get_container_client.return_value = cc
        blob_no_ct = self._blob("no_ct.bin", datetime(2026, 1, 1, tzinfo=UTC))
        blob_with_ct = self._blob(
            "data.csv", datetime(2026, 1, 2, tzinfo=UTC), content_type="text/csv"
        )
        cc.list_blobs.return_value = [blob_no_ct, blob_with_ct]
        cc.download_blob.return_value.readall.return_value = b"data"

        config = _config(
            "azure_blob",
            {"connection_string": "cs", "account_name": "acct", "container": "c1"},
        )
        with patch.dict("sys.modules", _fake_module_tree("azure.storage.blob", blob_mod)):
            docs = _run(_collect(AzureBlobConnector().get_delta(config, None)))

        assert len(docs) == 2
        by_name = {d.metadata["name"]: d for d, _ in docs}
        assert by_name["no_ct.bin"].content_type == "application/octet-stream"
        assert by_name["data.csv"].content_type == "text/csv"

    def test_skips_blob_older_than_cursor_and_excluded(self):
        from app.ingestion.connectors.azure_blob_connector import AzureBlobConnector

        blob_mod = MagicMock()
        service = MagicMock()
        blob_mod.BlobServiceClient.return_value = service
        cc = MagicMock()
        service.get_container_client.return_value = cc
        old_blob = self._blob("old.txt", datetime(2020, 1, 1, tzinfo=UTC))
        excluded_blob = self._blob("skip.tmp", datetime(2026, 1, 1, tzinfo=UTC))
        cc.list_blobs.return_value = [old_blob, excluded_blob]

        config = _config(
            "azure_blob", {"account_name": "a", "account_key": "k", "container": "c1"}
        )
        config.exclude_patterns = ["*.tmp"]
        with patch.dict("sys.modules", _fake_module_tree("azure.storage.blob", blob_mod)):
            docs = _run(_collect(
                AzureBlobConnector().get_delta(config, "2025-01-01T00:00:00+00:00")
            ))
        assert docs == []

    def test_download_error_skipped(self):
        from app.ingestion.connectors.azure_blob_connector import AzureBlobConnector

        blob_mod = MagicMock()
        service = MagicMock()
        blob_mod.BlobServiceClient.return_value = service
        cc = MagicMock()
        service.get_container_client.return_value = cc
        bad = self._blob("bad.txt", datetime(2026, 1, 1, tzinfo=UTC))
        cc.list_blobs.return_value = [bad]
        cc.download_blob.side_effect = RuntimeError("timeout")

        config = _config("azure_blob", {"account_name": "a", "account_key": "k", "container": "c1"})
        with patch.dict("sys.modules", _fake_module_tree("azure.storage.blob", blob_mod)):
            docs = _run(_collect(AzureBlobConnector().get_delta(config, None)))
        assert docs == []


# ═══════════════════════════════════════════════════════════════════════════
# SlackConnector
# ═══════════════════════════════════════════════════════════════════════════

class TestSlackConnectorValidateConnection:
    def test_success(self):
        from app.ingestion.connectors.slack_connector import SlackConnector

        config = _config("slack", {"bot_token": "xoxb-1"})
        with patch("httpx.AsyncClient") as mock_cls:
            resp = MagicMock()
            resp.json = MagicMock(return_value={"ok": True, "team": "Acme", "bot_id": "B1"})
            client = AsyncMock()
            client.get = AsyncMock(return_value=resp)
            client.__aenter__ = AsyncMock(return_value=client)
            client.__aexit__ = AsyncMock(return_value=False)
            mock_cls.return_value = client
            result = _run(SlackConnector().validate_connection(config))
        assert result.ok is True
        assert result.metadata["workspace"] == "Acme"

    def test_auth_failure(self):
        from app.ingestion.connectors.slack_connector import SlackConnector

        config = _config("slack", {"bot_token": "bad"})
        with patch("httpx.AsyncClient") as mock_cls:
            resp = MagicMock()
            resp.json = MagicMock(return_value={"ok": False, "error": "invalid_auth"})
            client = AsyncMock()
            client.get = AsyncMock(return_value=resp)
            client.__aenter__ = AsyncMock(return_value=client)
            client.__aexit__ = AsyncMock(return_value=False)
            mock_cls.return_value = client
            result = _run(SlackConnector().validate_connection(config))
        assert result.ok is False
        assert result.error == "invalid_auth"

    def test_exception_not_ok(self):
        from app.ingestion.connectors.slack_connector import SlackConnector

        config = _config("slack", {"bot_token": "x"})
        with patch("httpx.AsyncClient", side_effect=RuntimeError("dns error")):
            result = _run(SlackConnector().validate_connection(config))
        assert result.ok is False
        assert "dns error" in result.error


class TestSlackConnectorGetDelta:
    def test_yields_messages_as_documents(self):
        from app.ingestion.connectors.slack_connector import SlackConnector

        config = _config("slack", {"bot_token": "t", "channels": ["C1"]})
        chunks = [
            {
                "content": "hello world",
                "source_url": "https://slack.com/archives/C1",
                "metadata": {"ts": "1700000000.000100", "channel_id": "C1"},
            }
        ]
        with patch("app.knowledge.ingestors.slack_ingestor.SlackIngestor") as MockIngestor:
            MockIngestor.return_value.ingest_channel = AsyncMock(return_value=chunks)
            docs = _run(_collect(SlackConnector().get_delta(config, None)))

        assert len(docs) == 1
        doc, cursor = docs[0]
        assert doc.content == b"hello world"
        assert cursor == "1700000000.000100"

    def test_skips_messages_older_than_cursor(self):
        from app.ingestion.connectors.slack_connector import SlackConnector

        config = _config("slack", {"bot_token": "t", "channels": ["C1"]})
        chunks = [
            {"content": "old", "source_url": "u", "metadata": {"ts": "100.0"}},
            {"content": "new", "source_url": "u", "metadata": {"ts": "999.0"}},
        ]
        with patch("app.knowledge.ingestors.slack_ingestor.SlackIngestor") as MockIngestor:
            MockIngestor.return_value.ingest_channel = AsyncMock(return_value=chunks)
            docs = _run(_collect(SlackConnector().get_delta(config, "500.0")))

        assert len(docs) == 1
        assert docs[0][0].content == b"new"

    def test_channel_error_is_caught_and_continues(self):
        from app.ingestion.connectors.slack_connector import SlackConnector

        config = _config("slack", {"bot_token": "t", "channels": ["bad", "ok"]})
        good_chunks = [{"content": "fine", "source_url": "u", "metadata": {"ts": "1.0"}}]
        with patch("app.knowledge.ingestors.slack_ingestor.SlackIngestor") as MockIngestor:
            MockIngestor.return_value.ingest_channel = AsyncMock(
                side_effect=[RuntimeError("rate_limited"), good_chunks]
            )
            docs = _run(_collect(SlackConnector().get_delta(config, None)))
        assert len(docs) == 1
        assert docs[0][0].content == b"fine"

    def test_supports_streaming_true(self):
        from app.ingestion.connectors.slack_connector import SlackConnector

        assert SlackConnector().supports_streaming is True


# ═══════════════════════════════════════════════════════════════════════════
# PDFFileConnector / DOCXFileConnector
# ═══════════════════════════════════════════════════════════════════════════

def _mock_httpx_get(content=b"%PDF-1.4 fake", status_ok=True):
    resp = MagicMock()
    resp.content = content
    if status_ok:
        resp.raise_for_status = MagicMock()
    else:
        resp.raise_for_status = MagicMock(side_effect=RuntimeError("HTTP 404"))
    client = AsyncMock()
    client.get = AsyncMock(return_value=resp)
    client.__aenter__ = AsyncMock(return_value=client)
    client.__aexit__ = AsyncMock(return_value=False)
    return client


class TestPDFFileConnector:
    def test_validate_connection_always_ok(self):
        from app.ingestion.connectors.pdf_file_connector import PDFFileConnector

        result = _run(PDFFileConnector().validate_connection(_config("pdf_file")))
        assert result.ok is True

    def test_no_urls_yields_nothing(self):
        from app.ingestion.connectors.pdf_file_connector import PDFFileConnector

        config = _config("pdf_file", {})
        docs = _run(_collect(PDFFileConnector().get_delta(config, None)))
        assert docs == []

    def test_fetches_and_yields_document(self):
        from app.ingestion.connectors.pdf_file_connector import PDFFileConnector

        config = _config("pdf_file", {"urls": ["http://example.com/report.pdf"]})
        with patch("httpx.AsyncClient") as mock_cls:
            mock_cls.return_value = _mock_httpx_get(b"%PDF-1.4 content")
            docs = _run(_collect(PDFFileConnector().get_delta(config, None)))
        assert len(docs) == 1
        doc, cursor = docs[0]
        assert doc.content == b"%PDF-1.4 content"
        assert doc.content_type == "application/pdf"
        assert doc.title == "report.pdf"
        assert cursor == "http://example.com/report.pdf"

    def test_fetch_error_is_skipped(self):
        from app.ingestion.connectors.pdf_file_connector import PDFFileConnector

        config = _config(
            "pdf_file", {"urls": ["http://example.com/bad.pdf", "http://example.com/good.pdf"]}
        )
        good_client = _mock_httpx_get(b"good content")
        bad_client = _mock_httpx_get(status_ok=False)
        with patch("httpx.AsyncClient", side_effect=[bad_client, good_client]):
            docs = _run(_collect(PDFFileConnector().get_delta(config, None)))
        assert len(docs) == 1
        assert docs[0][0].content == b"good content"


class TestDOCXFileConnector:
    def test_validate_connection_always_ok(self):
        from app.ingestion.connectors.pdf_file_connector import DOCXFileConnector

        result = _run(DOCXFileConnector().validate_connection(_config("docx_file")))
        assert result.ok is True

    def test_fetches_and_yields_document(self):
        from app.ingestion.connectors.pdf_file_connector import DOCXFileConnector

        config = _config("docx_file", {"urls": ["http://example.com/doc.docx"]})
        with patch("httpx.AsyncClient") as mock_cls:
            mock_cls.return_value = _mock_httpx_get(b"docx bytes")
            docs = _run(_collect(DOCXFileConnector().get_delta(config, None)))
        assert len(docs) == 1
        doc, _cursor = docs[0]
        assert doc.content == b"docx bytes"
        assert "wordprocessingml" in doc.content_type
        assert doc.title == "doc.docx"

    def test_fetch_error_is_skipped(self):
        from app.ingestion.connectors.pdf_file_connector import DOCXFileConnector

        config = _config("docx_file", {"urls": ["http://example.com/bad.docx"]})
        with patch("httpx.AsyncClient") as mock_cls:
            mock_cls.return_value = _mock_httpx_get(status_ok=False)
            docs = _run(_collect(DOCXFileConnector().get_delta(config, None)))
        assert docs == []


# ═══════════════════════════════════════════════════════════════════════════
# ParquetParser
# ═══════════════════════════════════════════════════════════════════════════

class TestParquetParser:
    def test_parse_schema_and_rows(self):
        import pyarrow as pa
        import pyarrow.parquet as pq

        from app.ingestion.parsers.parquet_parser import ParquetParser

        table = pa.table({"id": [1, 2], "name": ["Alice", "Bob"]})
        buf = io.BytesIO()
        pq.write_table(table, buf)
        result = ParquetParser().parse(buf.getvalue(), filename="t.parquet")
        assert "Table: t.parquet" in result
        assert "Schema:" in result
        assert "id" in result and "name" in result
        assert "Row 0: id=1, name=Alice" in result

    def test_parse_without_filename_omits_table_header(self):
        import pyarrow as pa
        import pyarrow.parquet as pq

        from app.ingestion.parsers.parquet_parser import ParquetParser

        table = pa.table({"x": [1]})
        buf = io.BytesIO()
        pq.write_table(table, buf)
        result = ParquetParser().parse(buf.getvalue())
        assert "Table:" not in result

    def test_import_error_returns_empty(self):
        from app.ingestion.parsers.parquet_parser import ParquetParser

        with patch.dict("sys.modules", {"pyarrow": None}):
            result = ParquetParser().parse(b"anything", filename="x.parquet")
        assert result == ""

    def test_corrupt_bytes_returns_empty(self):
        from app.ingestion.parsers.parquet_parser import ParquetParser

        result = ParquetParser().parse(b"not a real parquet file", filename="bad.parquet")
        assert result == ""


# ═══════════════════════════════════════════════════════════════════════════
# AvroParser
# ═══════════════════════════════════════════════════════════════════════════

class TestAvroParser:
    def test_parse_schema_and_records(self):
        import fastavro

        from app.ingestion.parsers.avro_parser import AvroParser

        schema = {
            "type": "record",
            "name": "Person",
            "fields": [{"name": "id", "type": "int"}, {"name": "name", "type": "string"}],
        }
        records = [{"id": 1, "name": "Alice"}, {"id": 2, "name": "Bob"}]
        buf = io.BytesIO()
        fastavro.writer(buf, schema, records)
        result = AvroParser().parse(buf.getvalue(), filename="p.avro")
        assert "File: p.avro" in result
        assert "Schema: Person" in result
        assert "id (int)" in result
        assert "Record 0: id=1, name=Alice" in result

    def test_sample_records_limit(self):
        import fastavro

        from app.ingestion.parsers.avro_parser import AvroParser

        schema = {"type": "record", "name": "N", "fields": [{"name": "v", "type": "int"}]}
        records = [{"v": i} for i in range(60)]
        buf = io.BytesIO()
        fastavro.writer(buf, schema, records)
        parser = AvroParser()
        parser.SAMPLE_RECORDS = 5
        result = parser.parse(buf.getvalue())
        assert "Record 4:" in result
        assert "Record 5:" not in result

    def test_import_error_returns_empty(self):
        from app.ingestion.parsers.avro_parser import AvroParser

        with patch.dict("sys.modules", {"fastavro": None}):
            result = AvroParser().parse(b"anything", filename="x.avro")
        assert result == ""

    def test_corrupt_bytes_returns_empty(self):
        from app.ingestion.parsers.avro_parser import AvroParser

        result = AvroParser().parse(b"not an avro file", filename="bad.avro")
        assert result == ""


# ═══════════════════════════════════════════════════════════════════════════
# ExcelParser
# ═══════════════════════════════════════════════════════════════════════════

class TestExcelParser:
    def _make_workbook_bytes(self):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["Name", "Amount"])
        ws.append(["Alice", 100])
        ws.append([None, None])  # blank row, should be skipped
        ws.append(["Bob", 200])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_parse_sheet_rows_skip_blank(self):
        from app.ingestion.parsers.excel_parser import ExcelParser

        result = ExcelParser().parse(self._make_workbook_bytes(), filename="sales.xlsx")
        assert "Sheet: Sheet1" in result
        assert "Row: Name=Alice, Amount=100" in result
        assert "Row: Name=Bob, Amount=200" in result
        # exactly two data rows (blank row skipped)
        assert result.count("Row:") == 2

    def test_truncates_at_max_rows(self):
        import openpyxl

        from app.ingestion.parsers.excel_parser import ExcelParser

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["h1"])
        for i in range(10):
            ws.append([f"v{i}"])
        buf = io.BytesIO()
        wb.save(buf)

        parser = ExcelParser()
        parser.MAX_ROWS = 3
        result = parser.parse(buf.getvalue())
        assert f"[truncated at {parser.MAX_ROWS} rows]" in result

    def test_import_error_returns_empty(self):
        from app.ingestion.parsers.excel_parser import ExcelParser

        with patch.dict("sys.modules", {"openpyxl": None}):
            result = ExcelParser().parse(b"anything", filename="x.xlsx")
        assert result == ""

    def test_corrupt_bytes_returns_empty(self):
        from app.ingestion.parsers.excel_parser import ExcelParser

        result = ExcelParser().parse(b"not a real xlsx file", filename="bad.xlsx")
        assert result == ""

    def test_header_none_falls_back_to_col_index(self):
        import openpyxl

        from app.ingestion.parsers.excel_parser import ExcelParser

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([None, "B"])
        ws.append(["v1", "v2"])
        buf = io.BytesIO()
        wb.save(buf)

        result = ExcelParser().parse(buf.getvalue())
        assert "col0=v1" in result
        assert "B=v2" in result


# ── asyncio helper ───────────────────────────────────────────────────────────

def _run(coro):
    import asyncio

    return asyncio.run(coro)
