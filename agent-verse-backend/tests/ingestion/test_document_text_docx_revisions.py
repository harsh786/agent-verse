"""DOCX extraction reads the document as Word shows it with changes accepted (P1a-9).

Live P1a (KB-UPLOAD-HARD[docx-redlines]): python-docx's ``Paragraph.text`` only
reads runs that are direct children of the paragraph, so text inserted with
Track Changes (runs inside ``<w:ins>``) vanished ("Under clause 2.1, ; storage
...") and the amended demurrage period was never indexed. Headers and footers
(contract number, revision, approver) were not read at all. Deleted text must
stay out.
"""

from __future__ import annotations

import io
from typing import Any

from app.ingestion.document_text import extract_docx_text


def _tracked(kind: str, text: str, rid: int) -> Any:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    wrapper = OxmlElement(f"w:{kind}")
    wrapper.set(qn("w:id"), str(rid))
    wrapper.set(qn("w:author"), "Legal Ops")
    run = OxmlElement("w:r")
    t = OxmlElement("w:delText" if kind == "del" else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    wrapper.append(run)
    return wrapper


def _docx() -> bytes:
    from docx import Document

    d = Document()
    d.sections[0].header.paragraphs[0].text = "CONFIDENTIAL - Agreement HSA-4471"
    d.sections[0].footer.paragraphs[0].text = "Revision 7 approved by Amara Nwosu"
    d.add_heading("Demurrage", level=1)
    p = d.add_paragraph("Under clause 2.1, ")
    p._p.append(_tracked("del", "the free period is 4 days", 1))
    p._p.append(_tracked("ins", "the free period is 6 days", 2))
    p.add_run("; storage is billed per the tariff.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Service"
    table.cell(0, 1).text = "Response"
    table.cell(1, 0).text = "Reefer alarm"
    cell_p = table.cell(1, 1).paragraphs[0]
    cell_p._p.append(_tracked("del", "30 minutes", 3))
    cell_p._p.append(_tracked("ins", "20 minutes", 4))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_tracked_insertions_are_read_and_deletions_are_not() -> None:
    text = extract_docx_text(_docx())
    assert "Under clause 2.1, the free period is 6 days; storage is billed" in text
    assert "4 days" not in text


def test_tracked_changes_inside_table_cells_are_resolved_too() -> None:
    text = extract_docx_text(_docx())
    assert "Service: Reefer alarm; Response: 20 minutes" in text
    assert "30 minutes" not in text


def test_headers_and_footers_are_extracted_once() -> None:
    text = extract_docx_text(_docx())
    assert text.count("CONFIDENTIAL - Agreement HSA-4471") == 1
    assert text.count("Revision 7 approved by Amara Nwosu") == 1
    assert text.index("HSA-4471") < text.index("Demurrage") < text.index("Amara Nwosu")


def test_heading_is_a_markdown_heading() -> None:
    assert "# Demurrage" in extract_docx_text(_docx()).splitlines()
