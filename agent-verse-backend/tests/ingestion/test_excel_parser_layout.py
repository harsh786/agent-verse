"""ExcelParser understands real workbook layouts (P1a-10).

Live P1a (KB-UPLOAD-HARD[xlsx-merged]): a title row merged across the sheet
("Crane maintenance plan FY27") was taken as the header row, so every column
was named after the title or "col1".."col4" and the real header row became
data; a region merged down over its cranes reached only the first crane row,
so "STS-14" was indexed without "North Quay". Title rows are now captions, the
header row is the first row with at least two distinct labels, and a merged
range's value is applied to every cell it covers.
"""

from __future__ import annotations

import io

from app.ingestion.parsers.excel_parser import ExcelParser


def _workbook() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Cranes"
    ws.append(["Crane maintenance plan FY27 - Meridian Ports"])
    ws.merge_cells("A1:D1")
    ws.append(["Region", "Crane", "Next service", "Budget"])
    ws.append(["North Quay", "STS-11", "22 March 2027", 375000])
    ws.append([None, "STS-14", "12 February 2027", 775000])
    ws.merge_cells("A3:A4")
    ws.append(["South Basin", "STS-21", "13 March 2027", 750000])
    notes = wb.create_sheet("Notes")
    notes.append(["Note"])
    notes.append(["Spreader overhauls avoid the monsoon window."])
    grouped = wb.create_sheet("Quarterly")
    grouped.append(["Crane", "Moves", None, "Downtime"])
    grouped.merge_cells("B1:C1")
    grouped.append(["STS-11", 41000, 39000, "6 h"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _text() -> str:
    return ExcelParser().parse(_workbook(), filename="plan.xlsx")


def test_a_merged_title_row_is_a_caption_not_the_header() -> None:
    text = _text()
    assert "Title: Crane maintenance plan FY27 - Meridian Ports" in text
    assert "Row: Region=North Quay, Crane=STS-11, Next service=22 March 2027, Budget=375000" \
        in text
    assert "col1" not in text and "Crane maintenance plan FY27 - Meridian Ports=" not in text


def test_a_vertically_merged_cell_applies_to_every_row_it_covers() -> None:
    assert "Row: Region=North Quay, Crane=STS-14, Next service=12 February 2027" in _text()


def test_a_single_column_sheet_still_uses_its_first_row_as_header() -> None:
    assert "Row: Note=Spreader overhauls avoid the monsoon window." in _text()


def test_a_header_merged_across_columns_names_both_columns() -> None:
    assert "Row: Crane=STS-11, Moves=41000, Moves (2)=39000, Downtime=6 h" in _text()
