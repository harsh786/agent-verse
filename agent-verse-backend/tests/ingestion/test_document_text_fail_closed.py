"""PDF/DOCX uploads are parsed for real or refused — never indexed as raw bytes.

Regression: with pypdf absent (it was only an optional extra, and the production
image installs just ``--extra browser``) ``POST /knowledge/ingest/file`` decoded
the PDF's bytes as UTF-8 and indexed ``%PDF-1.3 … endobj``; ``/ingest/pdf``
embedded a placeholder chunk; both reported success. Parse failures were also
swallowed into ``[]`` (a 201 with 0 chunks). Uses real documents, no fakes.
"""

from __future__ import annotations

import io
from unittest.mock import patch

import pytest
from fastapi import HTTPException

from app.ingestion.document_text import (
    DocumentParseError,
    extract_docx_text,
    extract_pdf_pages,
)
from app.knowledge.ingestors.docx_ingestor import DocxIngestor
from app.knowledge.ingestors.pdf_ingestor import PdfIngestor


def _pdf(*pages: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    for text in pages:
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        if text:
            pdf.multi_cell(0, 6, text, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def _docx(*paragraphs: str) -> bytes:
    import docx

    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


_LEAVE = "Every full-time employee receives 22 days of paid annual leave per year."
_MEALS = "Meal expenses on business travel are reimbursed up to INR 1,500 per day."


def test_real_pdf_pages_are_extracted() -> None:
    pages = extract_pdf_pages(_pdf(_LEAVE, _MEALS))
    assert len(pages) == 2
    assert "22 days of paid annual leave" in pages[0]
    assert "INR 1,500 per day" in pages[1]


@pytest.mark.parametrize("data", [b"not a pdf at all", b"%PDF-1.4\n garbage"])
def test_unreadable_pdf_raises(data: bytes) -> None:
    with pytest.raises(DocumentParseError):
        extract_pdf_pages(data)


def test_pdf_without_text_layer_raises() -> None:
    with pytest.raises(DocumentParseError, match="no extractable text"):
        extract_pdf_pages(_pdf(""))


def test_pdf_ingestor_chunks_carry_page_citations() -> None:
    chunks = PdfIngestor().extract_chunks(content=_pdf(_LEAVE, _MEALS), filename="h.pdf")
    assert [c["page_number"] for c in chunks] == [1, 2]
    assert chunks[1]["metadata"] == {"filename": "h.pdf", "page": 2, "total_pages": 2}
    assert all("%PDF" not in c["content"] for c in chunks)


def test_pdf_ingestor_raises_instead_of_placeholder_or_empty() -> None:
    with pytest.raises(DocumentParseError):
        PdfIngestor().extract_chunks(content=b"garbage", filename="x.pdf")


def test_docx_text_and_ingestor() -> None:
    data = _docx("Leave policy", _LEAVE)
    assert _LEAVE in extract_docx_text(data)
    chunks = DocxIngestor().extract_chunks(content=data, filename="h.docx")
    assert chunks and _LEAVE in chunks[0]["content"]
    with pytest.raises(DocumentParseError):
        DocxIngestor().extract_chunks(content=b"PK\x03\x04 not a docx", filename="x.docx")


def test_pipeline_pdf_parser_never_decodes_bytes_as_text() -> None:
    from app.ingestion.parsers.pdf_parser import PDFParser

    parser = PDFParser()
    with (
        patch.object(PDFParser, "_parse_with_pymupdf", return_value=None),
        patch.object(PDFParser, "_parse_with_pdfminer", return_value=None),
    ):
        good = parser.parse_bytes(_pdf(_LEAVE), "h.pdf")
        bad = parser.parse_bytes(b"%PDF-1.4\n1 0 obj << >> endobj", "bad.pdf")
    assert "22 days" in good.full_text
    assert bad.full_text == "" and bad.error


def test_pipeline_docx_parser_refuses_without_library() -> None:
    from app.ingestion.parsers.docx_parser import DOCXParser

    data = _docx(_LEAVE)
    with patch.dict("sys.modules", {"docx": None}):
        result = DOCXParser().parse_bytes(data, "h.docx")
    assert result.paragraphs == [] and result.error


def test_upload_text_extraction_maps_to_http_errors() -> None:
    from app.api.knowledge import _extract_upload_text

    assert "22 days" in _extract_upload_text(_pdf(_LEAVE), ext="pdf", filename="h.pdf")
    assert _LEAVE in _extract_upload_text(_docx(_LEAVE), ext="docx", filename="h.docx")
    assert _extract_upload_text(b"plain notes", ext="txt", filename="n.txt") == "plain notes"
    for data, ext, code in (
        (b"garbage", "pdf", 422),
        (_pdf(""), "pdf", 422),
        (b"\xd0\xcf\x11\xe0 legacy word", "doc", 415),
        (b"\x89PNG\r\n\x1a\n\x00\x00", "txt", 415),
    ):
        with pytest.raises(HTTPException) as exc:
            _extract_upload_text(data, ext=ext, filename=f"f.{ext}")
        assert exc.value.status_code == code, (ext, exc.value.detail)
