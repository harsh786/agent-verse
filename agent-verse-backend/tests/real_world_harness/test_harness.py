"""Offline validation of the real-world suite's harness (no live stack needed).

The real-world scenarios only run against the live stack, so their fixtures,
scoring and reporting are checked here instead: the generated corpus against the
platform's own upload extractors and chunker (so expected chunk ranges and planted
facts are right before a live run ever depends on them), the workflow YAML against
the DSL model, publish checks and graph compiler, and the fixture server, metrics,
retrieval scoring and report builder as plain units.
"""

from __future__ import annotations

import io
import json
import os
import zipfile
from typing import Any

import httpx
import pytest
import yaml

from tests.real_world import corpus as cp
from tests.real_world import metrics as m
from tests.real_world import retrieval_eval as rev
from tests.real_world import wf_complex as wfc
from tests.real_world.fixture_server import FixtureServer, rss_feed

TEXT_FORMATS = ["pdf", "docx", "pptx", "xlsx", "csv", "html", "md"]


@pytest.fixture(scope="module")
def docs() -> dict[str, cp.CorpusDoc]:
    return {d.fmt: d for d in cp.build_corpus()}


def _extract(doc: cp.CorpusDoc) -> list[str]:
    """The platform's own upload extraction (what POST /knowledge/ingest/file does)."""
    from app.api.knowledge import _extract_upload_segments, _upload_ext

    segments, _pages = _extract_upload_segments(
        doc.data, ext=_upload_ext(doc.filename, doc.mime), filename=doc.filename)
    return [text for _page, text in segments]


def _chunks(texts: list[str]) -> int:
    """Chunks the upload endpoint makes of these segments (its structure-aware chunker)."""
    from app.knowledge.chunker_v2 import chunk_structured

    return sum(len([c for c, _ in chunk_structured(t, max_tokens=512, overlap_tokens=64)
                    if c.strip()]) for t in texts)


# ── corpus ─────────────────────────────────────────────────────────────────


def test_corpus_is_deterministic(docs: dict[str, cp.CorpusDoc]) -> None:
    assert len(docs) == 10
    # Formats without embedded timestamps are byte-identical on a rebuild.
    assert cp.build_csv(cp.doc_rng("csv")).data == docs["csv"].data
    assert cp.build_html(cp.doc_rng("html")).data == docs["html"].data
    assert cp.build_md(cp.doc_rng("md")).data == docs["md"].data
    assert cp.build_png().data == docs["png"].data
    # The others carry the same text.
    assert cp.build_docx(cp.doc_rng("docx")).text == docs["docx"].text


def test_pdf_shape(docs: dict[str, cp.CorpusDoc]) -> None:
    import pypdf

    pdf = docs["pdf"]
    assert 60 <= len(pypdf.PdfReader(io.BytesIO(pdf.data)).pages) == pdf.pages <= 120
    assert set(pdf.fact_pages) >= {*cp.PDF_FACTS, "sla_table", "sla_footnote"}


@pytest.mark.parametrize("fmt", TEXT_FORMATS)
def test_platform_extracts_and_chunks_within_range(fmt: str, docs: dict[str, cp.CorpusDoc]
                                                   ) -> None:
    doc = docs[fmt]
    texts = _extract(doc)
    joined = m.norm("\n".join(texts))
    n = _chunks(texts)
    lo, hi = doc.expected_chunks
    assert lo <= n <= hi, f"{doc.filename}: platform chunker gives {n}, fixture expects {lo}-{hi}"
    for q in cp.load_questions():
        if q["kind"] != "cross_doc" and m.source_matches(doc.filename, q["expected_sources"]):
            assert m.norm(q["chunk_must_contain"]) in joined, (q["id"], doc.filename)
            if q["kind"] not in ("table", "negation"):  # those answers are paraphrases
                assert any(m.norm(a) in joined for a in q["answer_any"]), (q["id"],
                                                                           doc.filename)
    for key, value in doc.table_probes:
        assert m.norm(key) in joined and m.norm(value) in joined, (doc.filename, key)


def test_pdf_pages_carry_their_facts(docs: dict[str, cp.CorpusDoc]) -> None:
    from app.ingestion.document_text import extract_pdf_pages

    pdf = docs["pdf"]
    pages = extract_pdf_pages(pdf.data, filename=pdf.filename)
    for q in cp.load_questions():
        if q.get("page_fact"):
            page = pdf.fact_pages[q["page_fact"]]
            assert m.norm(q["chunk_must_contain"]) in m.norm(pages[page - 1]), q["id"]


def test_scans_have_no_text_layer(docs: dict[str, cp.CorpusDoc]) -> None:
    import pypdf
    from PIL import Image

    reader = pypdf.PdfReader(io.BytesIO(docs["scan_pdf"].data))
    assert len(reader.pages) == 1 and not (reader.pages[0].extract_text() or "").strip()
    assert Image.open(io.BytesIO(docs["png"].data)).size == (1400, 900)


def test_xlsx_formulas_have_cached_values(docs: dict[str, cp.CorpusDoc]) -> None:
    from openpyxl import load_workbook

    data = load_workbook(io.BytesIO(docs["xlsx"].data), data_only=True)["Fuel"]
    formulas = load_workbook(io.BytesIO(docs["xlsx"].data))["Fuel"]
    assert str(formulas["C8"].value).startswith("=SUM(")
    assert data["C8"].value == cp.xlsx_fuel_total()
    assert {"Fleet", "Fuel", "Routes"} <= set(load_workbook(io.BytesIO(docs["xlsx"].data))
                                              .sheetnames)


def test_zip_members_hold_their_facts(docs: dict[str, cp.CorpusDoc]) -> None:
    z = zipfile.ZipFile(io.BytesIO(docs["zip"].data))
    assert set(z.namelist()) == set(cp.ZIP_FILES)
    assert "1.75 days" in z.read("leave-policy.txt").decode()


def test_docx_edit_changes_only_the_termination_clause() -> None:
    v1 = cp.build_docx(cp.doc_rng("docx")).text.splitlines()
    v2 = cp.build_docx_edited(cp.doc_rng("docx")).text.splitlines()
    changed = [(a, b) for a, b in zip(v1, v2, strict=True) if a != b]
    assert len(changed) == 1 and "120 days" in changed[0][1] and "75 days" in changed[0][0]
    assert "75 days" in cp.DOCX_FACTS["termination_notice"][1]  # global restored


def test_questions_are_well_formed() -> None:
    qs = cp.load_questions()
    assert len(qs) >= 25 and len({q["id"] for q in qs}) == len(qs)
    names = {d.filename for d in cp.build_corpus()} | set(cp.ZIP_FILES)
    kinds = {q["kind"] for q in qs}
    assert {"table", "cross_doc", "negation", "numeric", "ocr"} <= kinds
    for q in qs:
        assert set(q["expected_sources"]) <= names, q["id"]
        assert q["answer_any"] and q["chunk_must_contain"], q["id"]
    formula = next(q for q in qs if q["id"] == "xlsx-formula-total")
    assert str(cp.xlsx_fuel_total()) in formula["answer_any"]


def test_number_variants() -> None:
    assert cp._number_variants(23456789) == ["2,34,56,789", "23,456,789", "23456789"]
    assert cp._number_variants(950) == ["950"]


# ── workflows ──────────────────────────────────────────────────────────────


@pytest.mark.parametrize("builder", [
    lambda: wfc.complex_pipeline_yaml("rw-x", "https://fx.example", "k"),
    lambda: wfc.recovery_yaml("rw-r", "https://fx.example", "k"),
    lambda: wfc.chain_producer_yaml("rw-p", "rw-ch", "0123456789abcdef"),
    lambda: wfc.chain_consumer_yaml("rw-c", "rw-ch"),
    lambda: wfc.scheduled_ping_yaml("rw-s", "*/15 * * * *"),
])
def test_workflow_yaml_is_valid_publishable_and_compiles(builder: Any) -> None:
    from app.workflow.compiler import WorkflowCompiler
    from app.workflow.context import ContextResolver
    from app.workflow.dsl import WorkflowDefinition
    from app.workflow.service import publish_problems

    data = yaml.safe_load(builder())
    definition = WorkflowDefinition(**data)
    assert publish_problems(data) == []
    WorkflowCompiler(ContextResolver()).compile(definition)


def test_pipeline_shape_and_ground_truth() -> None:
    data = yaml.safe_load(wfc.complex_pipeline_yaml("rw-x", "https://fx.example", "k"))
    ids = [s["id"] for s in data["steps"]]
    assert ids == wfc.PIPELINE_TOP_LEVEL and len(ids) >= 10
    types = {s["type"] for s in data["steps"]}
    assert {"http", "code", "conditional", "parallel", "llm", "hitl", "set_variable"} <= types
    enrich = next(s for s in data["steps"] if s["id"] == "enrich")
    assert len(enrich["parallel_branches"]) == 3
    flaky = next(s for s in data["steps"] if s["id"] == "post_ledger")
    assert flaky["retry"]["max_attempts"] == 3
    exp = wfc.expected_pipeline()
    assert exp["lane"] == "expedite" and exp["gross"] == 82859.0 and exp["count"] == 12


def test_pipeline_code_steps_compute_ground_truth() -> None:
    """Run the parse/join code exactly as the sandbox would (inputs dict → output)."""
    from tests.real_world.fixture_server import CARRIER_SLA, FX_RATES, INVENTORY, load_orders

    ns: dict[str, Any] = {"inputs": {"orders": json.dumps(load_orders())}}
    exec(wfc.PARSE_ORDERS_CODE, ns)
    exp = wfc.expected_pipeline()
    assert {k: ns["output"][k] for k in ("count", "gross", "high_value", "cold_chain")} == {
        k: exp[k] for k in ("count", "gross", "high_value", "cold_chain")}
    ns = {"inputs": {"fx": FX_RATES, "inventory": INVENTORY, "sla": json.dumps(CARRIER_SLA),
                     "lane": "expedite", "gross": "82859.0"}}
    exec(wfc.JOIN_CODE, ns)
    assert ns["output"]["gross_usd"] == exp["gross_usd"]
    assert ns["output"]["best_carrier"] == exp["best_carrier"]


# ── fixture server ─────────────────────────────────────────────────────────


def test_fixture_server_routes_and_counters() -> None:
    with FixtureServer(port=0, host="127.0.0.1") as srv, httpx.Client(
            base_url=srv.local_base, timeout=10) as c:
        assert c.get("/health").text == "ok"
        assert len(c.get("/orders/k1.json").json()["orders"]) == 12
        codes = [c.post("/flaky/k1?fail=2", json={"x": 1}).status_code for _ in range(3)]
        assert codes == [503, 503, 200] and srv.count("POST", "/flaky/k1") == 3
        assert c.post("/charge/k1").json()["count"] == 1
        srv.set_switch("k1", True)
        assert c.get("/switch/k1").status_code == 500
        srv.set_switch("k1", False)
        assert c.get("/switch/k1").status_code == 200
        c.post("/publish/k1", json={"lane": "expedite"})
        assert srv.published["k1"] == [{"lane": "expedite"}]
        assert c.get("/feed/nope.xml").status_code == 404
        srv.set_feed("f", rss_feed([{"guid": "g1", "title": "T", "description": "D",
                                     "pubDate": "Mon, 28 Sep 2026 08:00:00 GMT"}]))
        import feedparser

        parsed = feedparser.parse(c.get("/feed/f.xml").text)
        assert [e.get("id") for e in parsed.entries] == ["g1"]


def test_fixture_public_base(monkeypatch: pytest.MonkeyPatch) -> None:
    srv = FixtureServer(port=4321)
    monkeypatch.delenv("RW_FIXTURE_PUBLIC_URL", raising=False)
    monkeypatch.setenv("RW_FIXTURE_HOST", "host.docker.internal")
    assert srv.public_base == "http://host.docker.internal:4321" and not srv.is_public()
    monkeypatch.setenv("RW_FIXTURE_PUBLIC_URL", "https://tunnel.example/")
    assert srv.public_base == "https://tunnel.example" and srv.is_public()


# ── metrics / scoring ──────────────────────────────────────────────────────


def test_metrics_basics() -> None:
    assert m.hit_at_k([1, 3, None, 6], 5) == 0.5
    assert m.mrr([1, 2, None, None]) == pytest.approx(0.375)
    assert m.percentiles([10, 20, 30, 40, 1000]) == {"n": 5, "mean": 220.0, "p50": 30.0,
                                                     "p95": 1000.0, "max": 1000.0}
    assert m.percentiles([]) == {"n": 0}
    assert m.norm("2 days – **Hosur**") == "2 days - hosur"
    assert m.source_matches("people-ops-bundle.zip/leave-policy.txt", ["leave-policy.txt"])
    assert not m.source_matches("larkctl-deploy-guide.md", ["it-disaster-recovery-runbook.html"])
    ev: dict[str, Any] = {}
    m.record(ev, a=0.123456, b=3)
    assert ev == {"metrics": {"a": 0.1235, "b": 3}}


def test_answer_correctness_rules() -> None:
    q = {"answer_any": ["not approved"], "answer_forbidden": ["drones are approved for last-mile"]}
    assert m.answer_correct("Drones are NOT approved in any zone.", q)
    assert not m.answer_correct("Yes - drones are approved for last-mile delivery.", q)
    cross = {"answer_any": ["75"], "answer_all": ["75", "45"]}
    assert m.answer_correct("Vendor: 75 days vs contractor: 45 days", cross)
    assert not m.answer_correct("Vendor: 75 days", cross)
    assert not m.answer_correct("", cross)


def test_question_rank_and_citations() -> None:
    q = {"kind": "table", "expected_sources": ["a.pdf"], "chunk_must_contain": "Nilgiri Zone"}
    hits = [{"source_file": "b.docx", "content": "Nilgiri Zone"},
            {"source_file": "a.pdf", "content": "x Nilgiri Zone 4 2", "page": "18"}]
    assert rev.question_rank(hits, q) == 2
    cross = {"kind": "cross_doc", "expected_sources": ["a.pdf", "b.docx"],
             "chunk_must_contain": "Nilgiri"}
    assert rev.question_rank(hits, cross) == 1
    assert rev.question_rank(hits[1:], cross) is None
    cites = [{"source": "a.pdf", "content": "Nilgiri Zone", "metadata": {"page": "18"}}]
    assert rev.citations_correct(cites, q, 18) and not rev.citations_correct(cites, q, 3)
    agg = rev.score([{"id": "1", "kind": "t", "rank": 1, "answered": True, "cited": True,
                      "search_ms": 10, "rag_ms": 100},
                     {"id": "2", "kind": "t", "rank": None, "answered": False, "cited": None,
                      "search_ms": 30, "rag_ms": 300}])
    assert agg["hit_at_5"] == 0.5 and agg["answer_accuracy"] == 0.5
    assert agg["citation_accuracy"] == 1.0 and agg["by_kind"]["t"]["n"] == 2


# ── report ─────────────────────────────────────────────────────────────────


def test_report_rollup_metrics_and_masking(tmp_path: Any) -> None:
    from tests.real_world import report

    rows = [
        {"suite": "backend", "scenario": "KB-COMPLEX-CORPUS", "test": "t::fmt[pdf]",
         "result": "passed", "duration_s": 1, "evidence": {"metrics": {"chunks": 90}},
         "failure_detail": ""},
        {"suite": "backend", "scenario": "KB-COMPLEX-CORPUS", "test": "t::fmt[zip]",
         "result": "failed", "duration_s": 1, "evidence": {},
         "failure_detail": "AssertionError: zip refused key av_abcdefghijklmnop123"},
        {"suite": "backend", "scenario": "KB-TENANT-ISOLATION", "test": "t::iso",
         "result": "skipped", "duration_s": 0, "evidence": {},
         "failure_detail": "Skipped: needs RW_SECOND_TENANT_API_KEY: a key of another tenant"},
        {"suite": "backend", "scenario": "KB-STRATEGIES", "test": "t::strat", "result": "passed",
         "duration_s": 3, "evidence": {"metrics": {"strategies": {
             "hybrid": {"available": True, "hit_at_5": 0.9, "p50_ms": 120},
             "colbert": {"available": False, "reason": "index_not_built", "probe_http": 503}}}},
         "failure_detail": ""},
    ]
    src = tmp_path / "rows.jsonl"
    src.write_text("\n".join(json.dumps(r) for r in rows))
    rep = report.build(str(src), "", str(tmp_path))
    roll = rep["scenarios"]
    assert roll["KB-COMPLEX-CORPUS"]["result"] == "failed"
    assert roll["KB-TENANT-ISOLATION"]["result"] == "skipped"
    assert "RW_SECOND_TENANT_API_KEY" in roll["KB-TENANT-ISOLATION"]["skip_reasons"][0]
    md = (tmp_path / "real_world_report.md").read_text()
    assert "av_abcdefghijklmnop123" not in md and "av_***" in md
    assert "| colbert | no |" in md and "chunks=90" in md
    assert json.loads((tmp_path / "real_world_report.json").read_text())["metrics"]


def test_runner_documents_every_env_var() -> None:
    """Each RW_* variable the suite reads is documented in the README."""
    import re

    here = os.path.join(os.path.dirname(__file__), "..", "real_world")
    used: set[str] = set()
    for name in os.listdir(here):
        if name.endswith(".py"):
            with open(os.path.join(here, name), encoding="utf-8") as fh:
                used |= set(re.findall(r"\b(RW_[A-Z0-9_]+)\b", fh.read()))
    with open(os.path.join(here, "README.md"), encoding="utf-8") as fh:
        readme = fh.read()
    missing = sorted(v for v in used if v not in readme)
    assert not missing, f"undocumented env vars: {missing}"
