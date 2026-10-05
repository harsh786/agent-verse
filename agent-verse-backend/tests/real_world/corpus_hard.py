"""Difficult upload cases for KB-UPLOAD-HARD (P1a): realistic, deterministic, generated.

The fictional owner is "Meridian Ports", a container-terminal operator. Every
fact is planted once, uses made-up proper nouns and has a known answer, so a
correct retrieval can only come from the right document (and, where it
matters, the right page / sheet / archive member).

Cases (``HardDoc.case``):
  pdf-2col, pdf-tables, pdf-owner-encrypted, pdf-cjk, pdf-mixed-scan, scan-3p,
  docx-redlines, pptx-notes, xlsx-merged, csv-large, csv-cp1252, html-boilerplate,
  md-code, docx-hindi, html-arabic, md-mixed-lang, png-rotated, png-lowq, zip-nested
and the refusals (``Refusal``): an encrypted (user password) PDF, a truncated PDF,
garbage named .docx, a zero-byte file, a truncated PNG, a zip bomb, a nested zip
bomb, and a file one byte over the upload limit. :func:`size_limit_docx` builds a
DOCX of exactly the limit.

Nothing here imports ``app``.
"""

from __future__ import annotations

import csv
import io
import os
import random
import zipfile
from dataclasses import dataclass, field
from typing import Any

from tests.real_world.corpus import (
    DOCX_MIME,
    PDF_MIME,
    PPTX_MIME,
    XLSX_MIME,
    _filler,
    _inject_cached_values,
    render_scan,
)

SEED = 20261005
UNICODE_FONT = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"


@dataclass
class HardQuestion:
    """A known-answer question about one document."""

    id: str
    question: str
    must_contain: str  # the hit's chunk must hold this (right chunk)
    answer_any: list[str]
    page: int | None = None  # 1-based PDF page the citation must name
    member: str | None = None  # archive member the hit must come from
    rag: bool = True  # also ask /rag/query (answer + citation)
    stale: str | None = None  # text that must NOT be served (superseded / deleted)


@dataclass
class HardDoc:
    case: str
    filename: str
    mime: str
    data: bytes
    expected_chunks: tuple[int, int]
    questions: list[HardQuestion] = field(default_factory=list)
    pages: int | None = None
    # substrings that must NOT appear in any chunk of this document (boilerplate,
    # script source, deleted revisions, ...)
    must_not_index: list[str] = field(default_factory=list)
    # substrings that must appear verbatim in some chunk (unicode / code fidelity)
    must_index: list[str] = field(default_factory=list)
    needs_ocr: bool = False
    max_upload_s: float = 60.0


@dataclass
class Refusal:
    case: str
    filename: str
    mime: str
    data: bytes
    statuses: tuple[int, ...]  # accepted HTTP statuses
    detail_any: list[str]  # the error must say why (any of these words)
    max_s: float = 30.0


def rng(case: str) -> random.Random:
    return random.Random(f"{SEED}:{case}")


def _has_unicode_font() -> bool:
    return os.path.exists(UNICODE_FONT)


# ── PDF ────────────────────────────────────────────────────────────────────────

PDF2COL_FACT = ("The Tamarind Bay terminal handled 1,284,500 TEU in FY26, a record for the "
                "Meridian Ports group, while the Saltmarsh terminal handled 402,300 TEU.")
PDF2COL_RIGHT_FACT = ("Quay crane productivity at Tamarind Bay averaged 31.4 moves per hour "
                      "across the year, up from 28.9 in FY25.")


def build_pdf_two_column() -> HardDoc:
    """A two-page annual-report spread set in two text columns (text flows down
    column 1, then column 2), with a full-width title above."""
    from fpdf import FPDF

    r = rng("pdf-2col")
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(True, margin=15)
    for page in (1, 2):
        pdf.add_page()
        pdf.set_font("Helvetica", style="B", size=15)
        pdf.cell(0, 9, f"Meridian Ports - Annual Report FY26 - Operations review ({page}/2)",
                 new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=9.5)
        body = [_filler(r, 140), PDF2COL_FACT if page == 1 else _filler(r, 60),
                _filler(r, 180), PDF2COL_RIGHT_FACT if page == 1 else _filler(r, 60),
                _filler(r, 120)]
        with pdf.text_columns(ncols=2, gutter=8, balance=True, text_align="J") as cols:
            for para in body:
                cols.write(text=para)
                cols.ln()
                cols.ln()
    return HardDoc(
        "pdf-2col", "meridian-annual-report-fy26.pdf", PDF_MIME, bytes(pdf.output()), (2, 8),
        pages=2,
        questions=[
            HardQuestion("2col-teu", "How many TEU did the Tamarind Bay terminal handle in FY26?",
                         "1,284,500 TEU in FY26", ["1,284,500", "1284500", "1.28 million"],
                         page=1),
            HardQuestion("2col-crane", "What was the average quay crane productivity at "
                         "Tamarind Bay in moves per hour?", "31.4 moves per hour", ["31.4"],
                         page=1, rag=False),
        ],
        # A column-interleaving extractor would split the sentence across chunks /
        # interleave the other column's words into it.
        must_index=["The Tamarind Bay terminal handled 1,284,500 TEU in FY26, a record for the"],
    )


TARIFF_ROWS = [
    ["Service", "Unit", "Rate (INR)", "Free period", "Notes"],
    ["Container handling 20ft", "per box", "6,450", "-", "laden"],
    ["Container handling 40ft", "per box", "9,800", "-", "laden"],
    ["Reefer plug-in", "per day", "2,150", "1 day", "includes monitoring"],
    ["Storage 20ft", "per day", "310", "5 days", "after free period"],
    ["Storage 40ft", "per day", "540", "5 days", "after free period"],
    ["Hazardous surcharge", "per box", "3,900", "-", "IMDG class 1-9"],
    ["Weighment (VGM)", "per box", "720", "-", "SOLAS"],
    ["Customs examination", "per box", "1,850", "-", "on request"],
    ["Out-of-gauge lift", "per lift", "18,600", "-", "pre-booked"],
    ["Shut-out cancellation", "per box", "4,200", "-", "within 24 hours"],
]
PENALTY_ROWS = [
    ["Delay band", "Penalty"],
    ["0-6 hours", "none"],
    ["6-24 hours", "INR 11,000 per vessel"],
    ["over 24 hours", "INR 46,000 per vessel"],
]


def build_pdf_tables() -> HardDoc:
    from fpdf import FPDF

    r = rng("pdf-tables")
    pdf = FPDF(format="A4")
    pdf.add_page()
    pdf.set_font("Helvetica", style="B", size=14)
    pdf.cell(0, 9, "Saltmarsh Terminal - Schedule of Tariffs 2027", new_x="LMARGIN",
             new_y="NEXT")
    pdf.set_font("Helvetica", size=9)
    pdf.multi_cell(0, 5, _filler(r, 80), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("Helvetica", size=8.5)
    with pdf.table(col_widths=(52, 22, 24, 24, 48), text_align="LEFT") as table:
        for row in TARIFF_ROWS:
            tr = table.row()
            for cell in row:
                tr.cell(cell)
    pdf.add_page()
    pdf.set_font("Helvetica", style="B", size=12)
    pdf.cell(0, 8, "Table 2 - Berth delay penalties", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9)
    with pdf.table(col_widths=(60, 80), text_align="LEFT") as table:
        for row in PENALTY_ROWS:
            tr = table.row()
            for cell in row:
                tr.cell(cell)
    pdf.multi_cell(0, 5, _filler(r, 120), new_x="LMARGIN", new_y="NEXT")
    return HardDoc(
        "pdf-tables", "saltmarsh-tariff-schedule-2027.pdf", PDF_MIME, bytes(pdf.output()),
        (2, 6), pages=2,
        questions=[
            HardQuestion("tariff-reefer", "What is the daily reefer plug-in rate at Saltmarsh "
                         "terminal?", "2,150", ["2,150", "2150"], page=1),
            HardQuestion("tariff-oog", "How much does an out-of-gauge lift cost?", "18,600",
                         ["18,600", "18600"], page=1, rag=False),
            HardQuestion("penalty-24h", "What is the berth delay penalty for delays over 24 "
                         "hours?", "46,000", ["46,000", "46000"], page=2),
        ],
        # The row must keep its label and its rate together.
        must_index=["Reefer plug-in", "Out-of-gauge lift"],
    )


OWNER_ENC_FACT = ("Pilotage for vessels above 300 metres LOA at Tamarind Bay requires two "
                  "pilots and the Heron tug pair.")


def _encrypted_pdf(user_password: str | None, text: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(0, 6, "Meridian Ports - Marine Services Circular MS-118",
                   new_x="LMARGIN", new_y="NEXT")
    pdf.multi_cell(0, 6, text, new_x="LMARGIN", new_y="NEXT")
    pdf.multi_cell(0, 6, _filler(rng("enc"), 90), new_x="LMARGIN", new_y="NEXT")
    pdf.set_encryption(owner_password="meridian-owner-2026", user_password=user_password)
    return bytes(pdf.output())


def build_pdf_owner_encrypted() -> HardDoc:
    """Permissions-only encryption (empty user password): every viewer opens it
    without a password, so the indexer must too (bank / government PDFs)."""
    return HardDoc(
        "pdf-owner-encrypted", "marine-circular-ms-118.pdf", PDF_MIME,
        _encrypted_pdf("", OWNER_ENC_FACT), (1, 2), pages=1,
        questions=[HardQuestion("enc-pilotage", "How many pilots are required for vessels "
                                "above 300 metres LOA at Tamarind Bay?", "two pilots",
                                ["two", "2 pilots"], page=1)],
    )


CJK_ZH = "青岚码头的冷藏集装箱插座共有 640 个，全部支持远程温度监控。"
CJK_JA = "白鷺ターミナルの営業時間は午前6時から午後10時までです。"


def build_pdf_cjk() -> HardDoc | None:
    if not _has_unicode_font():
        return None
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.add_font("ArialUni", fname=UNICODE_FONT)
    pdf.add_page()
    pdf.set_font("ArialUni", size=12)
    for line in ["Meridian Ports - 东亚码头指南 / 東アジア港湾ガイド", CJK_ZH,
                 "码头设有专用的危险品堆场，所有车辆进场前须完成安全检查。", CJK_JA,
                 "コンテナの搬入は事前予約制で、予約番号はゲートで確認されます。"]:
        pdf.multi_cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
    return HardDoc(
        "pdf-cjk", "east-asia-terminal-guide-cjk.pdf", PDF_MIME, bytes(pdf.output()), (1, 2),
        pages=1,
        questions=[
            HardQuestion("cjk-zh-reefer", "青岚码头有多少个冷藏集装箱插座？", "640",
                         ["640"], page=1),
            HardQuestion("cjk-ja-hours", "白鷺ターミナルの営業時間は？", "午前6時",
                         ["午前6時", "6時", "6:00", "6 am", "6am"], page=1, rag=False),
        ],
        must_index=[CJK_ZH, CJK_JA],
    )


MIXED_TYPED = ("Meridian Ports - Vessel incident letter VIL-3307. The berthing incident of "
               "the MV Calloway Star on 2 August 2026 is closed with no hull damage found.")
MIXED_SCAN_LINES = ["ANNEX A - SIGNED SURVEYOR NOTE", "Vessel: MV Calloway Star",
                    "Fender F-12 replaced by Okonkwo Marine", "Surveyor: Ingrid Halvorsen",
                    "Date: 05 August 2026"]


def _scan_page_pdf(images: list[Any], text_pages: list[str] | None = None) -> bytes:
    """A PDF whose pages are images (a scan); ``text_pages`` puts typed text pages first."""
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    for text in text_pages or []:
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        pdf.multi_cell(0, 6, text, new_x="LMARGIN", new_y="NEXT")
    for img in images:
        pdf.add_page()
        buf = io.BytesIO()
        img.convert("L").save(buf, format="PNG")
        buf.seek(0)
        pdf.image(buf, x=0, y=0, w=210, h=297)
    return bytes(pdf.output())


def build_pdf_mixed_scan() -> HardDoc:
    """Page 1 typed (text layer), page 2 a scanned annex (image only)."""
    scan = render_scan(MIXED_SCAN_LINES, seed=23, skew=0.4)
    data = _scan_page_pdf([scan], text_pages=[MIXED_TYPED + "\n\n" + _filler(rng("mix"), 80)])
    return HardDoc(
        "pdf-mixed-scan", "vessel-incident-letter-vil-3307.pdf", PDF_MIME, data, (2, 4),
        pages=2, needs_ocr=True, max_upload_s=120,
        questions=[
            HardQuestion("mix-typed", "Was hull damage found in the MV Calloway Star berthing "
                         "incident?", "no hull damage", ["no hull damage", "no damage"], page=1,
                         rag=False),
            HardQuestion("mix-scan", "Who replaced fender F-12 after the MV Calloway Star "
                         "incident?", "Okonkwo", ["okonkwo marine"], page=2),
        ],
    )


SCAN_PAGES = [
    ["MERIDIAN PORTS - CRANE INSPECTION REPORT", "Report no: CIR-2026-0912",
     "Terminal: Tamarind Bay", "Inspector: Rohan Deshpande", "Page 1 of 3"],
    ["SECTION 2 - FINDINGS", "Crane STS-09 hoist brake wear: 2.4 mm",
     "Wire rope replacement due: 14 December 2026", "Boom limit switch: OK", "Page 2 of 3"],
    ["SECTION 3 - SIGN OFF", "Crane STS-09 returned to service: 13 September 2026",
     "Approved by: Chief Engineer Lucia Ferreira", "Next inspection: March 2027",
     "Page 3 of 3"],
]


def build_scan_multipage() -> HardDoc:
    images = [render_scan(lines, seed=31 + i, skew=(-0.5, 0.7, -0.3)[i])
              for i, lines in enumerate(SCAN_PAGES)]
    return HardDoc(
        "scan-3p", "crane-inspection-report-cir-2026-0912-scan.pdf", PDF_MIME,
        _scan_page_pdf(images), (3, 6), pages=3, needs_ocr=True, max_upload_s=180,
        questions=[
            HardQuestion("scan3-brake", "What was the hoist brake wear measured on crane STS-09?",
                         "2.4 mm", ["2.4 mm", "2.4mm"], page=2),
            HardQuestion("scan3-approved", "Who approved crane STS-09's return to service?",
                         "Ferreira", ["lucia ferreira"], page=3),
            HardQuestion("scan3-inspector", "Who inspected the cranes for report CIR-2026-0912?",
                         "Deshpande", ["rohan deshpande"], page=1, rag=False),
        ],
    )


# ── DOCX ───────────────────────────────────────────────────────────────────────

DOCX_HEADER = "CONFIDENTIAL - Harbour Services Agreement HSA-4471 - Meridian Ports"
DOCX_FOOTER = "Document owner: Legal Operations, revision 7 (approved by Amara Nwosu)"
DOCX_INSERTED = "the demurrage free period is 6 calendar days from discharge"
DOCX_DELETED = "the demurrage free period is 4 calendar days from discharge"
DOCX_SLA_ROWS = [
    ["Region", "Service", "Response time", "Credit"],
    ["North Quay", "Crane breakdown", "45 minutes", "2% of monthly fee"],
    ["North Quay", "Reefer alarm", "20 minutes", "1% of monthly fee"],
    ["South Basin", "Crane breakdown", "60 minutes", "2% of monthly fee"],
    ["South Basin", "Gate outage", "30 minutes", "3% of monthly fee"],
]


def _tracked(kind: str, text: str, rid: int) -> Any:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    wrapper = OxmlElement(f"w:{kind}")
    wrapper.set(qn("w:id"), str(rid))
    wrapper.set(qn("w:author"), "Legal Ops")
    wrapper.set(qn("w:date"), "2026-09-01T10:00:00Z")
    run = OxmlElement("w:r")
    t = OxmlElement("w:delText" if kind == "del" else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    wrapper.append(run)
    return wrapper


def build_docx_redlines() -> HardDoc:
    """Header/footer, a table with vertically merged region cells, and tracked
    changes (one deletion, one insertion) inside a clause."""
    from docx import Document

    r = rng("docx-redlines")
    d = Document()
    section = d.sections[0]
    section.header.paragraphs[0].text = DOCX_HEADER
    section.footer.paragraphs[0].text = DOCX_FOOTER
    d.add_heading("Harbour Services Agreement", level=0)
    d.add_heading("1. Services", level=1)
    d.add_paragraph(_filler(r, 110))
    d.add_heading("2. Demurrage", level=1)
    p = d.add_paragraph("Under clause 2.1, ")
    p._p.append(_tracked("del", DOCX_DELETED, 1))
    p._p.append(_tracked("ins", DOCX_INSERTED, 2))
    p.add_run("; storage beyond the free period is billed per the tariff schedule.")
    d.add_paragraph(_filler(r, 110))
    d.add_heading("3. Service levels", level=1)
    table = d.add_table(rows=len(DOCX_SLA_ROWS), cols=4)
    table.style = "Table Grid"
    for i, row in enumerate(DOCX_SLA_ROWS):
        for j, cell in enumerate(row):
            if i > 0 and j == 0 and DOCX_SLA_ROWS[i - 1][0] == cell:
                continue  # merged below
            table.cell(i, j).text = cell
    table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(3, 0).merge(table.cell(4, 0))
    d.add_paragraph(_filler(r, 140))
    buf = io.BytesIO()
    d.save(buf)
    return HardDoc(
        "docx-redlines", "harbour-services-agreement-hsa-4471.docx", DOCX_MIME, buf.getvalue(),
        (1, 4),
        questions=[
            HardQuestion("docx-demurrage", "What is the demurrage free period under the Harbour "
                         "Services Agreement?", "6 calendar days", ["6 calendar days", "6 days",
                                                                   "six"],
                         stale="4 calendar days"),
            HardQuestion("docx-header", "Which agreement number is printed in the document "
                         "header?", "HSA-4471", ["hsa-4471"], rag=False),
            HardQuestion("docx-footer", "Who approved revision 7 of the agreement?",
                         "Amara Nwosu", ["amara nwosu"], rag=False),
            HardQuestion("docx-sla", "What is the reefer alarm response time at North Quay?",
                         "20 minutes", ["20 minutes"]),
        ],
        must_not_index=[DOCX_DELETED],
        must_index=[DOCX_INSERTED, "North Quay"],
    )


HINDI_FACT = "मेरिडियन बंदरगाह पर रात्रि पाली का भत्ता 850 रुपये प्रति पाली है।"


def build_docx_hindi() -> HardDoc:
    from docx import Document

    d = Document()
    d.add_heading("कर्मचारी परिपत्र संख्या 42", level=1)
    d.add_paragraph("यह परिपत्र सभी टर्मिनल कर्मचारियों के लिए है।")
    d.add_paragraph(HINDI_FACT)
    d.add_paragraph("सुरक्षा जूते और परावर्तक जैकेट हर समय पहनना अनिवार्य है।")
    d.add_paragraph("छुट्टी के आवेदन कम से कम सात दिन पहले जमा करें।")
    buf = io.BytesIO()
    d.save(buf)
    return HardDoc(
        "docx-hindi", "staff-circular-42-hindi.docx", DOCX_MIME, buf.getvalue(), (1, 2),
        questions=[HardQuestion("hi-allowance", "रात्रि पाली का भत्ता कितना है?", "850 रुपये",
                                ["850"])],
        must_index=[HINDI_FACT],
    )


# ── PPTX ───────────────────────────────────────────────────────────────────────

PPTX_NOTES_FACT = "Berth 7 dredging to 16.5 metres finishes on 30 March 2027"


def build_pptx_notes() -> HardDoc:
    from pptx import Presentation
    from pptx.util import Inches

    p = Presentation()
    slides = [
        ("Berth planning review", "Meridian Ports - Tamarind Bay, Q3 FY27",
         "Opening: thank the berth planners for the record quarter."),
        ("Berth utilisation", "Utilisation 71% (target 75%)",
         f"{PPTX_NOTES_FACT}; until then post-Panamax calls go to Berth 5."),
        ("Window compliance", None, "Compliance dipped in August because of monsoon swell."),
    ]
    for n, (title, body, notes) in enumerate(slides):
        s = p.slides.add_slide(p.slide_layouts[5])
        s.shapes.title.text = title
        if body:
            grp = s.shapes.add_group_shape()
            tb = grp.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(8), Inches(1))
            tb.text_frame.text = body
        if n == 2:
            rows = [["Line", "Window kept", "Calls"], ["Albatross Line", "88%", "41"],
                    ["Petrel Shipping", "64%", "29"]]
            shape = s.shapes.add_table(3, 3, Inches(0.7), Inches(1.8), Inches(8), Inches(1.5))
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    shape.table.cell(i, j).text = cell
        s.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    p.save(buf)
    return HardDoc(
        "pptx-notes", "berth-planning-review-q3.pptx", PPTX_MIME, buf.getvalue(), (1, 3),
        questions=[
            HardQuestion("pptx-dredging", "When does the Berth 7 dredging finish?",
                         "30 March 2027", ["30 march 2027", "march 30, 2027"]),
            HardQuestion("pptx-table", "What share of berth windows did Petrel Shipping keep?",
                         "64%", ["64%", "64 %"], rag=False),
        ],
        must_index=["Albatross Line | 88% | 41"],
    )


# ── XLSX ───────────────────────────────────────────────────────────────────────


def build_xlsx_merged() -> HardDoc:
    """A title row merged across the sheet, a header row below it, region cells
    merged vertically over their cranes, formulas with cached values, three sheets."""
    from openpyxl import Workbook

    r = rng("xlsx-merged")
    wb = Workbook()
    ws = wb.active
    ws.title = "Cranes"
    ws.append(["Crane maintenance plan FY27 - Meridian Ports"])
    ws.merge_cells("A1:E1")
    ws.append(["Region", "Crane", "Last service", "Next service", "Spares budget (INR)"])
    regions = [("North Quay", ["STS-11", "STS-12", "STS-14"]),
               ("South Basin", ["STS-21", "STS-22"]), ("East Mole", ["RTG-31", "RTG-32"])]
    row_no = 3
    budgets: list[int] = []
    for region, cranes in regions:
        start = row_no
        for crane in cranes:
            budget = 300_000 + 25_000 * r.randint(1, 20)
            nxt = "12 February 2027" if crane == "STS-14" else f"{r.randint(1, 28)} March 2027"
            ws.append([region if row_no == start else None, crane,
                       f"{r.randint(1, 28)} August 2026", nxt, budget])
            budgets.append(budget)
            row_no += 1
        if len(cranes) > 1:
            ws.merge_cells(f"A{start}:A{row_no - 1}")
    ws.append(["Total", None, None, None, f"=SUM(E3:E{row_no - 1})"])
    total_cell = f"E{row_no}"
    costs = wb.create_sheet("Contracts")
    costs.append(["Contractor", "Scope", "Annual value (INR)"])
    costs.append(["Okonkwo Marine", "Fenders and bollards", 2_450_000])
    costs.append(["Halvorsen Lifting", "Spreader overhaul", 3_975_000])
    costs.append(["Total", None, "=SUM(C2:C3)"])
    notes = wb.create_sheet("Notes")
    notes.append(["Note"])
    notes.append(["Spreader overhauls are scheduled outside the monsoon window."])
    buf = io.BytesIO()
    wb.save(buf)
    data = _inject_cached_values(buf.getvalue(), "xl/worksheets/sheet1.xml",
                                 {total_cell: sum(budgets)})
    data = _inject_cached_values(data, "xl/worksheets/sheet2.xml", {"C4": 6_425_000})
    total = sum(budgets)
    return HardDoc(
        "xlsx-merged", "crane-maintenance-plan-fy27.xlsx", XLSX_MIME, data, (1, 4),
        questions=[
            HardQuestion("xlsx-sts14", "When is crane STS-14 next serviced?", "STS-14",
                         ["12 february 2027"]),
            HardQuestion("xlsx-contracts-total", "What is the total annual value of the crane "
                         "maintenance contracts?", "6425000", ["6425000", "6,425,000",
                                                              "64,25,000", "6.4"]),
            HardQuestion("xlsx-budget-total", "What is the total spares budget?", str(total),
                         [str(total), f"{total:,}"], rag=False),
        ],
        # Headers must be the real header row, the merged region must reach every
        # crane row, and formula cells must carry their value (not be dropped).
        must_index=["Crane=STS-14", "Region=North Quay, Crane=STS-14"],
    )


# ── CSV ────────────────────────────────────────────────────────────────────────

CSV_LARGE_ROWS = int(os.getenv("RW_LARGE_CSV_ROWS", "20000"))
CSV_FACT = {"txn_id": "GT-2026-0177342", "gate": "Gate 4", "truck": "TN-38-BX-9921",
            "container": "MRDU 448812 7", "time": "2026-08-14 03:12",
            "remarks": "seal mismatch, sent to customs bay 2, cleared by \"R. Pillai\""}


def build_csv_large() -> HardDoc:
    r = rng("csv-large")
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=list(CSV_FACT), lineterminator="\n")
    w.writeheader()
    gates = ["Gate 1", "Gate 2", "Gate 3", "Gate 4"]
    remarks = ["ok", "ok", "ok", "late arrival, slot moved", "re-weighed, ok",
               "damaged door, photo taken\nsurveyor informed"]
    for i in range(1, CSV_LARGE_ROWS + 1):
        if i == CSV_LARGE_ROWS * 3 // 4:
            w.writerow(CSV_FACT)
            continue
        w.writerow({"txn_id": f"GT-2026-{100000 + i:07d}", "gate": r.choice(gates),
                    "truck": f"TN-{r.randint(10, 99)}-BX-{r.randint(1000, 9999)}",
                    "container": f"MRDU {r.randint(100000, 999999)} {r.randint(0, 9)}",
                    "time": f"2026-0{1 + i % 9}-{1 + i % 28:02d} {i % 24:02d}:{i % 60:02d}",
                    "remarks": r.choice(remarks)})
    data = buf.getvalue().encode()
    est = len(data) // 4 // 448
    return HardDoc(
        "csv-large", "gate-transactions-2026.csv", "text/csv", data,
        (max(1, est // 3), est * 4 + 2), max_upload_s=600,
        questions=[HardQuestion("csv-seal", "Which gate transaction had a seal mismatch for "
                                "truck TN-38-BX-9921?", "GT-2026-0177342",
                                ["gt-2026-0177342", "gate 4", "customs bay 2"])],
        # The quoted field with commas and quotes stays in its row.
        must_index=["seal mismatch, sent to customs bay 2"],
    )


CP1252_FACT = "Société Générale de Manutention handles lashing at Môle Est for €18,40 per box."


def build_csv_cp1252() -> HardDoc:
    rows = ["supplier;service;rate", "Société Générale de Manutention;lashing at Môle Est;"
            "€18,40 per box", "Brême Logistique;reefer pre-trip;€42,00 per unit"]
    data = ("\r\n".join(rows) + "\r\n").encode("cp1252")
    return HardDoc(
        "csv-cp1252", "fournisseurs-legacy-export.csv", "text/csv", data, (1, 1),
        questions=[HardQuestion("cp1252-lashing", "Who handles lashing at Môle Est?",
                                "Société Générale", ["société générale"], rag=False)],
        must_index=["Société Générale de Manutention", "€18,40"],
    )


# ── HTML ───────────────────────────────────────────────────────────────────────

HTML_FACT = ("Berth booking requests for the Kestrel Wharf close 72 hours before the vessel's "
             "estimated time of arrival.")
HTML_BOILERPLATE = ["Cookie preferences", "function trackVisitor", "Subscribe to our newsletter",
                    "All rights reserved Meridian Digital", "Related articles you may like",
                    "Skip to main content"]


def build_html_boilerplate() -> HardDoc:
    r = rng("html-boilerplate")
    nav = "".join(f'<li><a href="/p/{i}">Portal section {i}</a></li>' for i in range(1, 31))
    related = "".join(f"<li><a href='/a/{i}'>Story {i}: port news headline</a></li>"
                      for i in range(8))
    html = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<title>Kestrel Wharf berth booking - Meridian Ports portal</title>
<style>.banner{{position:fixed;bottom:0}} body{{font-family:sans-serif}}</style>
<script>function trackVisitor(id){{ window.dataLayer.push({{visitor:id, wharf:"Kestrel"}}); }}
trackVisitor("berth-booking-closes-24-hours");</script></head><body>
<a class="skip" href="#main">Skip to main content</a>
<header><div class="logo">Meridian Ports</div><nav><ul>{nav}</ul></nav></header>
<div class="banner" id="cookie">We use cookies. <button>Accept all</button>
<button>Cookie preferences</button></div>
<aside><h3>Related articles you may like</h3><ul>{related}</ul></aside>
<main id="main"><article>
<h1>Kestrel Wharf berth booking</h1>
<p>{HTML_FACT} Late requests are handled by the duty berth planner on a best-effort basis.</p>
<h2>How to book</h2>
<ol><li>Sign in to the berth window portal.</li><li>Upload the stowage plan (BAPLIE).</li>
<li>Confirm the pilot boarding time.</li></ol>
<p>{_filler(r, 120)}</p>
<table><tr><th>Wharf</th><th>Max draft</th></tr><tr><td>Kestrel Wharf</td><td>14.2 m</td></tr>
<tr><td>Osprey Wharf</td><td>12.8 m</td></tr></table>
</article></main>
<footer><p>Subscribe to our newsletter</p><p>&copy; 2026 All rights reserved Meridian Digital</p>
</footer></body></html>"""
    return HardDoc(
        "html-boilerplate", "kestrel-wharf-berth-booking.html", "text/html", html.encode(),
        (1, 2),
        questions=[
            HardQuestion("html-booking", "How long before arrival do Kestrel Wharf berth "
                         "booking requests close?", "72 hours", ["72 hours"],
                         stale="24-hours"),
            HardQuestion("html-draft", "What is the maximum draft at Kestrel Wharf?", "14.2",
                         ["14.2"], rag=False),
        ],
        must_not_index=HTML_BOILERPLATE,
    )


ARABIC_FACT = ("يجب على جميع العاملين في رصيف الياسمين ارتداء خوذة زرقاء اعتبارا من "
               "1 مارس 2027.")


def build_html_arabic() -> HardDoc:
    html = f"""<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8">
<title>تعميم السلامة رقم 9</title></head><body><main><article>
<h1>تعميم السلامة رقم 9 - موانئ ميريديان</h1>
<p>{ARABIC_FACT}</p>
<p>يمنع التدخين منعا باتا في منطقة الحاويات المبردة وفي جميع الأرصفة.</p>
<p>يجب الإبلاغ عن أي حادث خلال ساعة واحدة إلى مشرف الوردية المناوب.</p>
</article></main></body></html>"""
    return HardDoc(
        "html-arabic", "safety-circular-9-arabic.html", "text/html", html.encode(), (1, 1),
        questions=[HardQuestion("ar-helmet", "ما لون الخوذة المطلوبة في رصيف الياسمين؟",
                                "خوذة زرقاء", ["زرقاء", "blue"])],
        must_index=["رصيف الياسمين"],
    )


# ── Markdown ───────────────────────────────────────────────────────────────────

MD_CODE_LINES = [
    "def recalibrate_spreader(crane_id: str, twistlock_offset_mm: float) -> dict:",
    '    """Recalibrate the spreader twistlocks of one quay crane."""',
    "    if twistlock_offset_mm > 3.5:",
    '        raise ValueError("offset above 3.5 mm: take the crane out of service")',
]


def build_md_code() -> HardDoc:
    r = rng("md-code")
    code = "\n".join(MD_CODE_LINES + [f"    step_{i} = plc.write(crane_id, register={4000 + i})"
                                      for i in range(30)] + ["    return {'ok': True}"])
    md = (
        "# Crane telemetry runbook\n\n"
        "## Spreader recalibration\n\n"
        "Run the recalibration whenever twistlock offset alarms fire twice in one shift. "
        "The maximum tolerated twistlock offset is 3.5 mm.\n\n"
        f"```python\n{code}\n```\n\n"
        "## PLC registers\n\n| Register | Meaning |\n|---|---|\n| 4012 | hoist encoder |\n"
        "| 4017 | trolley encoder |\n\n"
        f"## Background\n\n{_filler(r, 200)}\n"
    )
    return HardDoc(
        "md-code", "crane-telemetry-runbook.md", "text/markdown", md.encode(), (1, 4),
        questions=[
            HardQuestion("md-offset", "Above which twistlock offset must the crane be taken "
                         "out of service?", "3.5 mm", ["3.5 mm", "3.5mm"]),
            HardQuestion("md-register", "Which PLC register is the trolley encoder?", "4017",
                         ["4017"], rag=False),
        ],
        # The function signature and its guard stay together (code not split mid-block).
        must_index=["\n".join(MD_CODE_LINES[:3])],
    )


MIXED_LANG = {
    "en": "The Meridian Ports night shift starts at 22:00 at every terminal.",
    "hi": "टर्मिनल पर ओवरटाइम का भुगतान हर महीने की 7 तारीख को होता है।",
    "ar": "يقع مكتب الجمارك الجديد بجوار البوابة رقم 3.",
    "zh": "所有危险品集装箱必须在进港前48小时申报。",
}


def build_md_mixed() -> HardDoc:
    md = "# Multilingual staff notice / बहुभाषी सूचना / إشعار / 通知\n\n" + "\n\n".join(
        f"## {lang}\n\n{text}" for lang, text in MIXED_LANG.items()) + "\n"
    return HardDoc(
        # Devanagari / Arabic cost several tokens per word: 1-3 small chunks.
        "md-mixed-lang", "multilingual-staff-notice.md", "text/markdown", md.encode(), (1, 3),
        questions=[
            HardQuestion("mixed-zh", "危险品集装箱必须提前多久申报？", "48小时", ["48"]),
            HardQuestion("mixed-hi", "ओवरटाइम का भुगतान कब होता है?", "7 तारीख", ["7"],
                         rag=False),
        ],
        must_index=list(MIXED_LANG.values()),
    )


# ── Images (OCR) ───────────────────────────────────────────────────────────────

ROTATED_LINES = ["MERIDIAN PORTS GATE PASS", "Pass no: GP-77120", "Driver: Tomasz Wieczorek",
                 "Valid until: 31 October 2026"]
LOWQ_LINES = ["RECEIPT - HERON TUGS", "Towage job: TJ-5531", "Amount: INR 1,86,000",
              "Paid by: Albatross Line"]


def build_png_rotated() -> HardDoc:
    """A phone photo of a gate pass taken sideways (page rotated 90 degrees)."""
    img = render_scan(ROTATED_LINES, size=(1400, 900), font_px=52, skew=0.3, seed=41)
    img = img.rotate(90, expand=True)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return HardDoc("png-rotated", "gate-pass-gp-77120-sideways.png", "image/png",
                   buf.getvalue(), (1, 1), needs_ocr=True,
                   questions=[HardQuestion("rot-driver", "Who is the driver on gate pass "
                                           "GP-77120?", "Wieczorek", ["tomasz wieczorek"])])


def build_png_lowq() -> HardDoc:
    """A low-resolution, blurred, noisy, JPEG-damaged receipt photo saved as PNG."""
    from PIL import Image, ImageFilter

    img = render_scan(LOWQ_LINES, size=(1400, 900), font_px=54, skew=-1.2, seed=43)
    img = img.resize((700, 450), Image.Resampling.BILINEAR).filter(
        ImageFilter.GaussianBlur(0.8))
    r = rng("png-lowq")
    px = img.load()
    for _ in range(700 * 450 // 40):
        x, y = r.randrange(700), r.randrange(450)
        px[x, y] = max(0, min(255, px[x, y] + r.randint(-70, 70)))
    jpg = io.BytesIO()
    img.save(jpg, format="JPEG", quality=30)
    jpg.seek(0)
    out = io.BytesIO()
    Image.open(jpg).save(out, format="PNG")
    return HardDoc("png-lowq", "heron-tugs-receipt-tj-5531-photo.png", "image/png",
                   out.getvalue(), (1, 1), needs_ocr=True,
                   questions=[HardQuestion("lowq-job", "Which towage job number is on the "
                                           "Heron Tugs receipt?", "TJ-5531", ["tj-5531"])])


# ── ZIP ────────────────────────────────────────────────────────────────────────

ZIP_NAME = "terminal-handbook-bundle.zip"
ZIP_OVERTIME = "Overtime at Meridian Ports is paid at 1.75 times the base hourly rate."
ZIP_HISTORY = "The Saltmarsh terminal was commissioned in 1987 by the Brackwater Harbour Trust."
ZIP_CONTACT = ("Duty harbour master", "Captain Ines Valdivia", "+91 44 5550 1187")


def _docx_bytes(paragraphs: list[str]) -> bytes:
    from docx import Document

    d = Document()
    for para in paragraphs:
        d.add_paragraph(para)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _small_pdf(lines: list[str]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in lines:
        pdf.multi_cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def build_zip_nested() -> HardDoc:
    """Mixed formats, a nested archive two levels down, macOS junk, an unsupported
    legacy file and a path-traversal member name."""
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("notes/history.md", f"# Terminal history\n\n{ZIP_HISTORY}\n")
        z.writestr("notes/.DS_Store", b"\x00\x00\x00\x01Bud1" + bytes(64))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("handbook/overtime-policy.docx", _docx_bytes(
            ["Overtime policy", ZIP_OVERTIME, "Overtime must be pre-approved by the shift lead."]))
        z.writestr("handbook/tariff-extract.pdf", _small_pdf(
            ["Tariff extract", "Pilot cancellation within 2 hours of boarding: INR 27,500."]))
        z.writestr("contacts.csv", "role,name,phone\n" + ",".join(ZIP_CONTACT) + "\n"
                   "Gate supervisor,Arjun Mehta,+91 44 5550 1190\n")
        z.writestr("archive/archive-2025.zip", inner.getvalue())
        z.writestr("__MACOSX/handbook/._overtime-policy.docx", b"\x00\x05\x16\x07" + bytes(80))
        z.writestr("legacy/old-roster.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(512))
        z.writestr("../../outside/escape-note.txt",
                   "Escape note: the Brackwater ferry runs every 40 minutes.\n")
    return HardDoc(
        "zip-nested", ZIP_NAME, "application/zip", buf.getvalue(), (4, 10),
        questions=[
            HardQuestion("zip-overtime", "At what multiple of the base rate is overtime paid?",
                         "1.75 times", ["1.75"], member="overtime-policy.docx"),
            HardQuestion("zip-history", "Who commissioned the Saltmarsh terminal in 1987?",
                         "Brackwater Harbour Trust", ["brackwater harbour trust"],
                         member="history.md"),
            HardQuestion("zip-contact", "Who is the duty harbour master?", "Valdivia",
                         ["ines valdivia"], member="contacts.csv", rag=False),
            HardQuestion("zip-pdf", "What is the pilot cancellation fee within 2 hours of "
                         "boarding?", "27,500", ["27,500", "27500"], member="tariff-extract.pdf",
                         rag=False),
        ],
        must_not_index=["Bud1"],
    )


def zip_bomb(member_mib: int = 1024) -> bytes:
    """One deflated member of ``member_mib`` MiB of zeros (~1000:1)."""
    buf = io.BytesIO()
    block = bytes(1024 * 1024)
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z, z.open("payload.txt", "w") as fh:
        for _ in range(member_mib):
            fh.write(block)
    return buf.getvalue()


def nested_zip_bomb() -> bytes:
    """Ten copies of a 256 MiB bomb inside an outer archive (each small on its own)."""
    inner = zip_bomb(256)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as z:
        for i in range(10):
            z.writestr(f"part-{i}.zip", inner)
    return buf.getvalue()


# ── Size limit ─────────────────────────────────────────────────────────────────

SIZE_FACT = "The Plover breakwater extension adds 640 metres of sheltered quay."


def size_limit_docx(total_bytes: int) -> bytes:
    """A DOCX of exactly ``total_bytes``: a short text plus an incompressible
    embedded media part (like a large scanned figure)."""
    base = _docx_bytes(["Breakwater extension brief", SIZE_FACT])

    def with_pad(pad: int) -> bytes:
        src = zipfile.ZipFile(io.BytesIO(base))
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                dst.writestr(item, src.read(item.filename))
            dst.writestr(zipfile.ZipInfo("word/media/figure-survey.bin"),
                         random.Random(SEED).randbytes(pad), compress_type=zipfile.ZIP_STORED)
        return out.getvalue()

    pad = max(0, total_bytes - len(base) - 200)
    data = with_pad(pad)
    data = with_pad(pad + (total_bytes - len(data)))
    if len(data) != total_bytes:
        raise AssertionError(f"size_limit_docx: {len(data)} != {total_bytes}")
    return data


# ── Duplicates ─────────────────────────────────────────────────────────────────

SHARED_PARAGRAPH = ("The Kingfisher gate closes at 23:00 for maintenance every Sunday; trucks "
                    "are diverted to Gate 2 during the closure.")


def duplicate_docs() -> dict[str, bytes]:
    a = (f"# Ops bulletin 31\n\n{SHARED_PARAGRAPH}\n\nYard block C7 reopens on Monday.\n")
    b = (f"# Ops bulletin 32\n\n{SHARED_PARAGRAPH}\n\nReefer row R4 is closed for "
         "electrical repairs until Friday.\n")
    return {"ops-bulletin-31.md": a.encode(), "ops-bulletin-31-copy.md": a.encode(),
            "ops-bulletin-31.txt": a.encode(), "ops-bulletin-32.md": b.encode()}


# ── The corpus ─────────────────────────────────────────────────────────────────


def build_hard_docs() -> list[HardDoc]:
    docs = [build_pdf_two_column(), build_pdf_tables(), build_pdf_owner_encrypted(),
            build_pdf_cjk(), build_pdf_mixed_scan(), build_scan_multipage(),
            build_docx_redlines(), build_docx_hindi(), build_pptx_notes(), build_xlsx_merged(),
            build_csv_large(), build_csv_cp1252(), build_html_boilerplate(), build_html_arabic(),
            build_md_code(), build_md_mixed(), build_png_rotated(), build_png_lowq(),
            build_zip_nested()]
    return [d for d in docs if d is not None]


def build_refusals() -> list[Refusal]:
    good_pdf = build_pdf_tables().data
    png = build_png_lowq().data
    return [
        Refusal("pdf-encrypted", "board-minutes-restricted.pdf", PDF_MIME,
                _encrypted_pdf("board-only-7731", "Restricted minutes."), (422,),
                ["encrypted", "password"]),
        Refusal("pdf-truncated", "tariff-schedule-truncated.pdf", PDF_MIME,
                good_pdf[: len(good_pdf) * 2 // 5], (422,), ["pdf", "readable", "corrupt"]),
        Refusal("docx-garbage", "not-really-a-document.docx", DOCX_MIME,
                random.Random(SEED).randbytes(40_000), (422,), ["docx", "readable"]),
        Refusal("empty-file", "empty-notes.txt", "text/plain", b"", (422,), ["empty"]),
        Refusal("png-truncated", "half-uploaded-photo.png", "image/png", png[: len(png) // 3],
                (422,), ["image", "readable"]),
        Refusal("zip-bomb", "quarterly-exports.zip", "application/zip", zip_bomb(), (413, 422),
                ["archive", "zip", "uncompressed", "ratio", "bomb"]),
        Refusal("zip-bomb-nested", "nested-exports.zip", "application/zip", nested_zip_bomb(),
                (413, 422), ["archive", "zip", "uncompressed", "ratio", "bomb"]),
    ]
