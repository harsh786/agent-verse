"""A programmable web site the live stack crawls (P1d: URL ingest and web crawl).

The site is ``fixture_server.py`` running in a throwaway container on the compose
network (``agentverse-rw-web``, aliases ``rw-web`` and ``rw-web-b``, labelled
``p1d=live-test``), operator-allowlisted for the ingestion egress guard. The stack
fetches ``http://rw-web:8080/...``; this process registers pages and reads the
request log through the published control port (``RW_WEB_CONTROL_URL``, default
``http://127.0.0.1:58080``)::

    docker run -d --name agentverse-rw-web -l p1d=live-test \\
      --network agentverse-backend_default --network-alias rw-web --network-alias rw-web-b \\
      -p 127.0.0.1:58080:8080 -v $PWD/tests/real_world:/srv:ro \\
      python:3.12-slim python -u /srv/fixture_server.py --serve --port 8080

Every scenario works under its own ``/site/<key>/`` prefix; ``/robots.txt`` is one
per host, so a scenario that needs robots rules sets them for its prefix.
"""

from __future__ import annotations

import base64
import io
import os
from typing import Any

import httpx
import pytest

from tests.real_world.helpers import tag


def control_url() -> str:
    return os.getenv("RW_WEB_CONTROL_URL", "http://127.0.0.1:58080").rstrip("/")


def public_host() -> str:
    """How the stack reaches the site (an operator-allowlisted compose alias)."""
    return os.getenv("RW_WEB_HOST", "rw-web:8080")


def other_host() -> str:
    """A second name of the same server: another site, for crawl scoping."""
    return os.getenv("RW_WEB_OTHER_HOST", "rw-web-b:8080")


class WebSite:
    """One scenario's pages under ``/site/<key>/`` on the fixture web server."""

    def __init__(self, key: str | None = None) -> None:
        self.key = key or tag()
        self.prefix = f"/site/{self.key}"
        self.ctl = httpx.Client(base_url=control_url(), timeout=60)

    @classmethod
    def open(cls) -> WebSite:
        site = cls()
        try:
            site.ctl.get("/health").raise_for_status()
        except httpx.HTTPError as exc:
            pytest.skip(f"needs the web fixture container ({control_url()}): {exc}")
        return site

    def close(self) -> None:
        self.ctl.close()

    # ── pages ────────────────────────────────────────────────────────────────
    def path(self, rel: str) -> str:
        return f"{self.prefix}/{rel.lstrip('/')}" if rel else f"{self.prefix}/"

    def url(self, rel: str = "", host: str | None = None) -> str:
        return f"http://{host or public_host()}{self.path(rel)}"

    def put(self, rel: str, body: str | bytes = b"", *, status: int = 200,
            content_type: str | None = "text/html; charset=utf-8",
            headers: dict[str, str] | None = None, absolute: bool = False,
            **spec: Any) -> None:
        """Serve ``body`` at ``rel`` (under the prefix, or as-is when ``absolute``)."""
        data = body.encode("utf-8") if isinstance(body, str) else body
        all_headers = dict(headers or {})
        if content_type is not None:
            all_headers.setdefault("Content-Type", content_type)
        path = rel if absolute else self.path(rel)
        resp = self.ctl.put("/_control/route", json={
            "path": path, "status": status, "headers": all_headers,
            "body_b64": base64.b64encode(data).decode(), **spec})
        resp.raise_for_status()

    def redirect(self, rel: str, location: str, status: int = 301) -> None:
        self.put(rel, b"", status=status, content_type=None, headers={"Location": location})

    def remove(self, rel: str) -> None:
        self.ctl.delete("/_control/route", params={"path": self.path(rel)}).raise_for_status()

    def robots(self, text: str) -> None:
        self.put("/robots.txt", text, content_type="text/plain", absolute=True)

    def hits(self) -> list[dict[str, Any]]:
        return list(self.ctl.get("/_control/hits", params={"prefix": self.prefix}).json())

    def all_hits(self, prefix: str = "/") -> list[dict[str, Any]]:
        return list(self.ctl.get("/_control/hits", params={"prefix": prefix}).json())


def page(title: str, *paragraphs: str, links: tuple[tuple[str, str], ...] = (),
         head: str = "", chrome: bool = True, lang: str = "en") -> str:
    """A realistic help-center page: header nav, cookie banner, sidebar, footer,
    inline CSS and JS around a ``<main>`` article."""
    body = "".join(f"<p>{p}</p>" for p in paragraphs)
    related = "".join(f'<li><a href="{href}">{text}</a></li>' for href, text in links)
    if not chrome:
        return (f"<!doctype html><html lang='{lang}'><head><title>{title}</title>{head}</head>"
                f"<body><main><h1>{title}</h1>{body}<ul>{related}</ul></main></body></html>")
    return f"""<!doctype html><html lang="{lang}"><head><meta name="viewport" content="width=device-width">
<title>{title} | Kestrel Logistics Help Center</title>{head}
<style>body{{font-family:Inter,sans-serif}} .cookie-banner{{position:fixed;bottom:0}}</style>
<script>window.dataLayer=window.dataLayer||[];function gtag(){{dataLayer.push(arguments)}}
gtag('config','G-KESTREL01');</script></head><body>
<a class="skip-link" href="#main">Skip to main content</a>
<header class="masthead"><nav aria-label="Primary"><ul><li><a href="/">Kestrel home</a></li>
<li><a href="/careers">Careers at Kestrel</a></li><li><a href="/investors">Investor relations</a></li>
</ul></nav></header>
<div class="cookie-banner" id="cookie-consent">We use cookies to personalise content and ads.
<button>Accept all cookies</button></div>
<aside class="sidebar"><h3>Popular articles</h3><ul><li>How to reset your portal password</li>
<li>Track a shipment</li></ul></aside>
<main id="main"><article><h1>{title}</h1>{body}
<h2>Related</h2><ul>{related}</ul></article></main>
<footer><p>© 2026 Kestrel Logistics Pvt Ltd · Subscribe to our newsletter</p>
<script>trackFooterImpression('kb');</script></footer></body></html>"""


CHROME = ("Careers at Kestrel", "Investor relations", "Accept all cookies",
          "We use cookies", "Popular articles", "Subscribe to our newsletter", "dataLayer",
          "trackFooterImpression", "font-family", "Skip to main content")


def pdf_bytes(pages: list[list[str]]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    for lines in pages:
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        for line in lines:
            pdf.multi_cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def docx_bytes(title: str, paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading(title, level=1)
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
