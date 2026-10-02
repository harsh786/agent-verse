"""A realistic, deterministic multi-format knowledge corpus for the complex KB scenarios.

The corpus is the document set of a fictional logistics company ("Larkspur
Logistics"): a 60-120 page operations policy manual (PDF with sections, tables and
footnotes), a vendor master agreement (DOCX with headings and tables), a quarterly
review deck (PPTX with speaker notes), a fleet workbook (multi-sheet XLSX with
formulas and cached values), a ~5,000-row shipment ledger (CSV), an IT runbook
(HTML with nested lists), a deploy guide (Markdown with code blocks), a scanned
delivery note (image-only PDF, the OCR path), a whiteboard photo (PNG, OCR) and an
HR bundle (ZIP of mixed files).

Every fact the scenarios ask about is planted exactly once, uses made-up proper
nouns (so a correct answer can only come from the knowledge base), and is listed
in ``fixtures/kb_questions.json``. The filler text around the facts is generated
from a fixed seed, so the same corpus (same bytes, same chunk counts) is produced
on every run. Nothing here imports ``app``; the offline harness tests
(``tests/real_world_harness``) check the corpus against the platform's own
extractors and chunker.
"""

from __future__ import annotations

import csv
import io
import json
import math
import os
import random
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any

FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures")

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Chunker parameters of POST /knowledge/ingest/file (chunk_by_tokens 512 / 64).
CHUNK_TOKENS = 512
CHUNK_OVERLAP = 64


@dataclass
class CorpusDoc:
    """One generated document plus what the scenarios expect from ingesting it."""

    filename: str
    fmt: str  # pdf | docx | pptx | xlsx | csv | html | md | scan_pdf | png | zip
    mime: str
    data: bytes
    text: str  # the text an extractor should recover (for the offline checks)
    expected_chunks: tuple[int, int]  # inclusive range of chunks_created
    headings: list[str] = field(default_factory=list)  # section headings in the doc
    # (heading, fact substring) pairs: the chunk holding the fact should carry the heading
    heading_probes: list[tuple[str, str]] = field(default_factory=list)
    # (row key, value) pairs: a table row that must stay intact inside one chunk
    table_probes: list[tuple[str, str]] = field(default_factory=list)
    pages: int | None = None
    fact_pages: dict[str, int] = field(default_factory=dict)  # fact id -> 1-based PDF page
    needs_ocr: bool = False


# ── Planted facts ──────────────────────────────────────────────────────────────
# (ids match fixtures/kb_questions.json)

PDF_NAME = "larkspur-ops-policy-manual.pdf"
DOCX_NAME = "bramblewood-vendor-master-agreement.docx"
PPTX_NAME = "q3-fy27-operations-review.pptx"
XLSX_NAME = "fleet-and-fuel-fy27.xlsx"
CSV_NAME = "shipment-ledger-2026.csv"
HTML_NAME = "it-disaster-recovery-runbook.html"
MD_NAME = "larkctl-deploy-guide.md"
SCAN_NAME = "delivery-note-dn-58213-scan.pdf"
PNG_NAME = "wh9-whiteboard-photo.png"
ZIP_NAME = "people-ops-bundle.zip"

PDF_SECTIONS = [
    "1. Purpose and scope", "2. Roles and responsibilities", "3. Network and zones",
    "4. Regional service levels", "5. Packaging standards", "6. Hazardous goods",
    "7. Cold-chain handling", "8. Customs and documentation", "9. Incident management",
    "10. Contractors and exits", "11. Last-mile delivery", "12. Records retention",
]

PDF_FACTS: dict[str, tuple[str, str]] = {
    # fact id: (section, sentence)
    "cold_chain_range": (
        "7. Cold-chain handling",
        "Pharmaceutical consignments must stay between 2 deg C and 8 deg C at all times; "
        "any temperature excursion longer than 25 minutes requires a Form CC-19 deviation "
        "report signed by the shift lead.",
    ),
    "incident_escalation": (
        "9. Incident management",
        "Any incident with an estimated loss above INR 4,00,000 is escalated to the Regional "
        "Risk Council within 6 hours of detection.",
    ),
    "contractor_exit_notice": (
        "10. Contractors and exits",
        "A contractor exit requires a written notice period of 45 days, counted from the "
        "date the Larkspur exit memo is acknowledged.",
    ),
    "drones_not_approved": (
        "11. Last-mile delivery",
        "Drones are NOT approved for last-mile delivery in any Larkspur zone as of policy "
        "version 9.2; pilot requests must go to the Innovation Board.",
    ),
    "retention_period": (
        "12. Records retention",
        "Proof-of-delivery records are retained for 8 years in the Juniper archive before "
        "certified destruction.",
    ),
    "hazmat_class": (
        "6. Hazardous goods",
        "Lithium battery shipments are handled as UN3480 Class 9 and may only be loaded "
        "on vehicles fitted with a Type-C fire blanket.",
    ),
}

# Regional SLA table (section 4) — table-lookup questions.
SLA_TABLE = [
    ["Zone", "Standard (days)", "Express (days)", "Late penalty"],
    ["Nilgiri Zone", "4", "2", "3.5%"],
    ["Deccan Zone", "3", "1", "2.0%"],
    ["Konkan Zone [1]", "5", "3", "4.0%"],
    ["Vindhya Zone", "6", "3", "2.5%"],
    ["Brahmaputra Zone", "7", "4", "5.0%"],
]
SLA_FOOTNOTE = (
    "[1] The Konkan Zone late penalty is waived during the monsoon months of June to "
    "September; standard and express targets still apply."
)

DOCX_SECTIONS = ["Definitions", "Scope of services", "Fees and invoicing", "Termination",
                 "Liability", "Data protection", "Vendor register"]
DOCX_FACTS = {
    "termination_notice": (
        "Termination",
        "Either party may terminate this agreement for convenience with a written notice "
        "period of 75 days.",
    ),
    "liability_cap": (
        "Liability",
        "Bramblewood Freight's aggregate liability is capped at 1.5 times the annual fees "
        "paid in the preceding twelve months.",
    ),
    "invoice_terms": (
        "Fees and invoicing",
        "Invoices are payable within 40 days of receipt; late payments accrue interest at "
        "1.25% per month.",
    ),
}
DOCX_AMENDED_TERMINATION = (
    "Either party may terminate this agreement for convenience with a written notice "
    "period of 120 days (Amendment No. 3)."
)
VENDOR_TABLE = [
    ["Vendor", "Contract ID", "Renewal date", "Annual value"],
    ["Bramblewood Freight", "VMA-2207", "31 January 2027", "INR 62 lakh"],
    ["Corvid Cold Storage", "VMA-2231", "15 April 2027", "INR 38 lakh"],
    ["Halyard Customs Brokers", "VMA-2245", "30 June 2027", "INR 21 lakh"],
]

PPTX_SLIDES = [
    ("Q3 FY27 operations review", "Larkspur Logistics - leadership review, October 2026",
     "Welcome everyone; this deck covers network performance for July to September."),
    ("On-time performance", "OTP 94.2% against a 95% target; express OTP 97.1%",
     "The Hosur hub caused most of the Q3 delays after its conveyor outage on 12 August."),
    ("Cost per shipment", "Cost per shipment fell to INR 182 from INR 197 in Q2",
     "Savings came mainly from route consolidation in the Deccan Zone."),
    ("Claims", "Damage claims: 311 (Q2: 402); average settlement 9 days",
     "Claims dropped after the new edge protectors were rolled out."),
    ("People", "Attrition 3.1% quarterly; 46 new drivers onboarded",
     "Driver onboarding now takes 6 days instead of 11."),
    ("Q4 priorities", "Peak season readiness; Vindhya cross-dock go-live 18 November",
     "The Vindhya cross-dock go-live date is fixed for 18 November; no slips allowed."),
]

FLEET_FACT_VEHICLE = ("KA-51-LX-7719", "Eicher Pro 2049", 3200)
FUEL_MONTHS = ["April", "May", "June", "July", "August", "September"]

CSV_FACT_ROW = {
    "shipment_id": "SHP-2026-04417", "date": "2026-07-19", "origin": "Coimbatore",
    "destination": "Guwahati", "weight_kg": "1840", "carrier": "Kestrel Cargo",
    "status": "held at customs",
}

HTML_FACTS = {
    "tier1": ("Tier 1 systems", "RPO 15 minutes, RTO 1 hour"),
    "tier2": ("Tier 2 systems", "RPO 4 hours, RTO 12 hours"),
    "tier3": ("Tier 3 systems", "RPO 24 hours, RTO 72 hours"),
}

MD_CANARY_CMD = "larkctl rollout --canary 10 --region ap-south-1"

SCAN_LINES = [
    "LARKSPUR LOGISTICS - DELIVERY NOTE",
    "Delivery note: DN-58213",
    "Consignee: Corvid Cold Storage, Hosur",
    "Pallets delivered: 14",
    "Received by: Meenakshi Iyer",
    "Date: 03 September 2026",
]
PNG_LINES = [
    "WH-9 FIRE DRILL",
    "Every 3rd Friday at 16:30",
    "Assembly point: Gate B",
    "Warden: Farhan Qureshi",
]

ZIP_FILES = {
    "leave-policy.txt": (
        "People Ops leave policy (Larkspur Logistics).\n"
        "Employees accrue 1.75 days of earned leave per month of service.\n"
        "Unused earned leave above 45 days lapses on 31 March each year.\n"
    ),
    "onboarding.md": (
        "# Onboarding checklist\n\n"
        "- Laptop requests go to the Saffron desk on the 3rd floor.\n"
        "- Badge photos are taken on Tuesdays only.\n"
    ),
    "holidays-2027.csv": (
        "date,holiday\n2027-01-26,Republic Day\n2027-03-04,Holi\n2027-08-15,Independence Day\n"
    ),
}


# ── Filler text ────────────────────────────────────────────────────────────────

_SUBJECTS = ["The shift lead", "Each hub manager", "The operations desk", "Every driver",
             "The quality team", "The regional planner", "Warehouse supervisors",
             "The compliance officer", "Dispatch coordinators", "The safety committee"]
_VERBS = ["reviews", "records", "verifies", "reconciles", "documents", "audits",
          "schedules", "inspects", "approves", "monitors"]
_OBJECTS = ["the daily load plan", "outbound manifests", "dock assignments",
            "vehicle checklists", "exception reports", "returns paperwork",
            "pallet labels", "seal numbers", "handover logs", "yard movements"]
_TAILS = ["before the first wave departs.", "at the end of every shift.",
          "within the agreed handover window.", "using the standard template.",
          "and escalates gaps to the duty manager.", "as described in the local SOP.",
          "with a second person present.", "and keeps the evidence for audit.",
          "before any exception is closed.", "in line with the zone calendar."]


def _filler(rng: random.Random, words: int) -> str:
    out: list[str] = []
    count = 0
    while count < words:
        s = (f"{rng.choice(_SUBJECTS)} {rng.choice(_VERBS)} {rng.choice(_OBJECTS)} "
             f"{rng.choice(_TAILS)}")
        out.append(s)
        count += len(s.split())
    return " ".join(out)


def _est_chunks(text: str) -> int:
    """Rough chunk estimate for one segment (~4 chars per token, 448-token stride)."""
    tokens = max(1, len(text) // 4)
    return max(1, math.ceil((tokens - CHUNK_OVERLAP) / (CHUNK_TOKENS - CHUNK_OVERLAP)))


def _range(est: int, lo: float = 0.5, hi: float = 2.5) -> tuple[int, int]:
    return max(1, math.floor(est * lo)), math.ceil(est * hi) + 2


# ── Format builders ────────────────────────────────────────────────────────────


def build_pdf(rng: random.Random, pages: int) -> CorpusDoc:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(True, margin=18)
    pdf.set_title("Larkspur Logistics - Operations Policy Manual v9.2")
    section_pages = {s: 1 + (i * (pages - 1)) // len(PDF_SECTIONS)
                     for i, s in enumerate(PDF_SECTIONS)}
    fact_by_section: dict[str, list[tuple[str, str]]] = {}
    for fid, (sec, sentence) in PDF_FACTS.items():
        fact_by_section.setdefault(sec, []).append((fid, sentence))
    texts: list[str] = []
    fact_pages: dict[str, int] = {}
    current = PDF_SECTIONS[0]
    for page in range(1, pages + 1):
        pdf.add_page()
        page_text: list[str] = []
        pdf.set_font("Helvetica", size=8)
        header = f"Larkspur Logistics - Operations Policy Manual v9.2 - page {page} of {pages}"
        pdf.cell(0, 5, header, new_x="LMARGIN", new_y="NEXT")
        page_text.append(header)
        starting = [s for s, p in section_pages.items() if p == page]
        if starting:
            current = starting[0]
            pdf.set_font("Helvetica", style="B", size=14)
            pdf.multi_cell(0, 8, current, new_x="LMARGIN", new_y="NEXT")
            page_text.append(current)
        else:
            pdf.set_font("Helvetica", style="B", size=11)
            sub = f"{current} (continued)"
            pdf.multi_cell(0, 6, sub, new_x="LMARGIN", new_y="NEXT")
            page_text.append(sub)
        pdf.set_font("Helvetica", size=10)
        if starting:
            for fid, sentence in fact_by_section.get(current, []):
                pdf.multi_cell(0, 5, sentence, new_x="LMARGIN", new_y="NEXT")
                page_text.append(sentence)
                fact_pages[fid] = page
        if starting and current == "4. Regional service levels":
            intro = "Table 4.1 - Regional service levels by zone:"
            pdf.multi_cell(0, 5, intro, new_x="LMARGIN", new_y="NEXT")
            page_text.append(intro)
            with pdf.table(col_widths=(50, 40, 40, 40), text_align="LEFT") as table:
                for row in SLA_TABLE:
                    r = table.row()
                    for cell in row:
                        r.cell(cell)
            page_text.extend(" ".join(r) for r in SLA_TABLE)
            fact_pages["sla_table"] = page
        body = _filler(rng, 170 if starting else 230)
        pdf.multi_cell(0, 5, body, new_x="LMARGIN", new_y="NEXT")
        page_text.append(body)
        if starting and current == "4. Regional service levels":
            pdf.set_font("Helvetica", size=7)
            pdf.multi_cell(0, 4, SLA_FOOTNOTE, new_x="LMARGIN", new_y="NEXT")
            page_text.append(SLA_FOOTNOTE)
            fact_pages["sla_footnote"] = page
        texts.append("\n".join(page_text))
    data = bytes(pdf.output())
    est = sum(_est_chunks(t) for t in texts)
    return CorpusDoc(
        PDF_NAME, "pdf", PDF_MIME, data, "\n".join(texts), _range(est, 0.7, 2.0),
        headings=list(PDF_SECTIONS),
        heading_probes=[(sec, sentence.split(";")[0][:60])
                        for sec, sentence in PDF_FACTS.values()],
        table_probes=[(row[0].replace(" [1]", ""), row[2]) for row in SLA_TABLE[1:]],
        pages=pages, fact_pages=fact_pages,
    )


def build_docx(rng: random.Random) -> CorpusDoc:
    from docx import Document

    d = Document()
    d.add_heading("Vendor Master Agreement - Bramblewood Freight (VMA-2207)", level=0)
    texts: list[str] = ["Vendor Master Agreement - Bramblewood Freight (VMA-2207)"]
    for sec in DOCX_SECTIONS:
        d.add_heading(sec, level=1)
        texts.append(sec)
        for fid, (fsec, sentence) in DOCX_FACTS.items():
            if fsec == sec:
                d.add_paragraph(sentence)
                texts.append(sentence)
        if sec == "Vendor register":
            table = d.add_table(rows=len(VENDOR_TABLE), cols=len(VENDOR_TABLE[0]))
            table.style = "Table Grid"
            for i, row in enumerate(VENDOR_TABLE):
                for j, cell in enumerate(row):
                    table.cell(i, j).text = cell
            texts.extend(" ".join(r) for r in VENDOR_TABLE)
        for _ in range(3):
            para = _filler(rng, 120)
            d.add_paragraph(para)
            texts.append(para)
    buf = io.BytesIO()
    d.save(buf)
    text = "\n".join(texts)
    return CorpusDoc(
        DOCX_NAME, "docx", DOCX_MIME, buf.getvalue(), text, _range(_est_chunks(text)),
        headings=list(DOCX_SECTIONS),
        heading_probes=[(sec, sentence[:60]) for sec, sentence in DOCX_FACTS.values()],
        table_probes=[(row[0], row[1]) for row in VENDOR_TABLE[1:]],
    )


def build_docx_edited(rng: random.Random) -> CorpusDoc:
    """The vendor agreement after an amendment: only the Termination clause changes.

    Pass an rng in the same state as for the original (see :func:`doc_rng`) so every
    other paragraph is byte-identical.
    """
    original = DOCX_FACTS["termination_notice"]
    DOCX_FACTS["termination_notice"] = ("Termination", DOCX_AMENDED_TERMINATION)
    try:
        doc = build_docx(rng)
    finally:
        DOCX_FACTS["termination_notice"] = original
    return doc


def build_pptx(rng: random.Random) -> CorpusDoc:
    from pptx import Presentation
    from pptx.util import Inches

    p = Presentation()
    texts: list[str] = []
    for title, body, notes in PPTX_SLIDES:
        s = p.slides.add_slide(p.slide_layouts[5])
        s.shapes.title.text = title
        tb = s.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(8.5), Inches(3))
        tb.text_frame.text = body
        s.notes_slide.notes_text_frame.text = notes
        texts += [title, body, f"Notes: {notes}"]
    buf = io.BytesIO()
    p.save(buf)
    text = "\n".join(texts)
    return CorpusDoc(PPTX_NAME, "pptx", PPTX_MIME, buf.getvalue(), text,
                     _range(_est_chunks(text), 0.5, 3.0),
                     headings=[t for t, _, _ in PPTX_SLIDES])


def _inject_cached_values(xlsx: bytes, sheet_xml: str, values: dict[str, float]) -> bytes:
    """Give formula cells a cached value, as Excel does when it saves a workbook.

    openpyxl writes formulas without results; a real workbook carries both the
    formula and its last computed value, which is what an indexer reads.
    """
    src = zipfile.ZipFile(io.BytesIO(xlsx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            raw = src.read(item.filename)
            if item.filename == sheet_xml:
                xml = raw.decode()
                for cell, value in values.items():
                    xml = re.sub(
                        rf'(<c r="{cell}"[^>]*>)(<f>[^<]*</f>)(<v\s*/>|<v></v>)?',
                        lambda m, v=value: f"{m.group(1)}{m.group(2)}<v>{v}</v>",
                        xml,
                    )
                raw = xml.encode()
            dst.writestr(item, raw)
    return out.getvalue()


def build_xlsx(rng: random.Random) -> CorpusDoc:
    from openpyxl import Workbook

    wb = Workbook()
    fleet = wb.active
    fleet.title = "Fleet"
    fleet.append(["vehicle_id", "model", "capacity_kg", "home_hub"])
    texts = ["Fleet vehicle_id model capacity_kg home_hub"]
    models = ["Tata Ultra 1918", "Ashok Leyland Boss 1415", "BharatBenz 1617R",
              "Mahindra Furio 14"]
    hubs = ["Hosur", "Coimbatore", "Pune", "Nagpur", "Guwahati"]
    for i in range(38):
        row = [f"KA-{10 + i:02d}-LX-{1000 + 37 * i}", rng.choice(models),
               rng.choice([1800, 2500, 4000, 6000]), rng.choice(hubs)]
        if i == 17:
            row = [FLEET_FACT_VEHICLE[0], FLEET_FACT_VEHICLE[1], FLEET_FACT_VEHICLE[2], "Hosur"]
        fleet.append(row)
        texts.append(" ".join(str(c) for c in row))
    fuel = wb.create_sheet("Fuel")
    fuel.append(["month", "diesel_litres", "diesel_cost_inr", "cng_cost_inr"])
    litres = [41200 + 1350 * i for i in range(6)]
    costs = [round(v * 93.4) for v in litres]
    cng = [182000 + 4500 * i for i in range(6)]
    for m, li, co, cn in zip(FUEL_MONTHS, litres, costs, cng, strict=True):
        fuel.append([m, li, co, cn])
        texts.append(f"{m} {li} {co} {cn}")
    last = len(FUEL_MONTHS) + 1
    fuel.append(["Total H1", f"=SUM(B2:B{last})", f"=SUM(C2:C{last})", f"=SUM(D2:D{last})"])
    totals = {f"B{last + 1}": sum(litres), f"C{last + 1}": sum(costs),
              f"D{last + 1}": sum(cng)}
    texts.append(f"Total H1 {sum(litres)} {sum(costs)} {sum(cng)}")
    routes = wb.create_sheet("Routes")
    routes.append(["route", "distance_km", "avg_hours"])
    for i, (a, b) in enumerate([("Hosur", "Pune"), ("Pune", "Nagpur"), ("Nagpur", "Guwahati"),
                                ("Coimbatore", "Hosur")]):
        row = [f"{a}-{b}", 300 + 211 * i, 6 + 3 * i]
        routes.append(row)
        texts.append(" ".join(str(c) for c in row))
    buf = io.BytesIO()
    wb.save(buf)
    data = _inject_cached_values(buf.getvalue(), "xl/worksheets/sheet2.xml", totals)
    text = "\n".join(texts)
    return CorpusDoc(XLSX_NAME, "xlsx", XLSX_MIME, data, text,
                     _range(_est_chunks(text), 0.4, 4.0),
                     headings=["Fleet", "Fuel", "Routes"],
                     table_probes=[(FLEET_FACT_VEHICLE[0], str(FLEET_FACT_VEHICLE[2]))])


def xlsx_fuel_total() -> int:
    litres = [41200 + 1350 * i for i in range(6)]
    return sum(round(v * 93.4) for v in litres)


def build_csv(rng: random.Random, rows: int = 5000) -> CorpusDoc:
    cities = ["Hosur", "Pune", "Nagpur", "Guwahati", "Coimbatore", "Indore", "Surat",
              "Vizag", "Kochi", "Ranchi"]
    carriers = ["Larkspur Fleet", "Osprey Movers", "Tern Express", "Plover Roadlines"]
    statuses = ["delivered", "in transit", "out for delivery", "delivered", "delivered"]
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=list(CSV_FACT_ROW), lineterminator="\n")
    w.writeheader()
    special = 4417
    for i in range(1, rows + 1):
        if i == special:
            w.writerow(CSV_FACT_ROW)
            continue
        a, b = rng.sample(cities, 2)
        w.writerow({
            "shipment_id": f"SHP-2026-{i:05d}", "date": f"2026-{1 + i % 9:02d}-{1 + i % 27:02d}",
            "origin": a, "destination": b, "weight_kg": str(rng.randint(40, 9000)),
            "carrier": rng.choice(carriers), "status": rng.choice(statuses),
        })
    text = buf.getvalue()
    return CorpusDoc(CSV_NAME, "csv", "text/csv", text.encode(), text,
                     _range(_est_chunks(text), 0.3, 4.0),
                     table_probes=[(CSV_FACT_ROW["shipment_id"], CSV_FACT_ROW["status"])])


def build_html(rng: random.Random) -> CorpusDoc:
    tiers = "".join(
        f"<li>{name}<ul><li>Recovery objectives: {objective}</li>"
        f"<li>Owner: {owner}</li></ul></li>"
        for (name, objective), owner in zip(HTML_FACTS.values(),
                                            ["Platform SRE", "Data Engineering",
                                             "Business Apps"], strict=True)
    )
    filler = "".join(f"<p>{_filler(rng, 90)}</p>" for _ in range(6))
    html = (
        "<!doctype html><html><head><title>IT disaster recovery runbook</title></head><body>"
        "<h1>IT disaster recovery runbook</h1>"
        "<h2>Recovery tiers</h2>"
        f"<ul>{tiers}</ul>"
        "<h2>Failover procedure</h2>"
        "<ol><li>Declare the incident in the #dr-bridge channel</li>"
        "<li>Fail over DNS to the Chennai standby region<ul>"
        "<li>Run <code>larkctl dns failover --to maa</code></li>"
        "<li>Confirm health checks are green for 10 minutes</li></ul></li>"
        "<li>Notify the Regional Risk Council</li></ol>"
        f"<h2>Background</h2>{filler}"
        "</body></html>"
    )
    text = re.sub(r"<[^>]+>", " ", html)
    return CorpusDoc(HTML_NAME, "html", "text/html", html.encode(), text,
                     _range(_est_chunks(text), 0.5, 3.0),
                     headings=["Recovery tiers", "Failover procedure", "Background"],
                     heading_probes=[("Tier 2 systems", "RPO 4 hours")])


def build_md(rng: random.Random) -> CorpusDoc:
    md = (
        "# larkctl deploy guide\n\n"
        "## Prerequisites\n\n"
        "- VPN connected to the Larkspur build network\n"
        "- `larkctl` 4.2 or later on your PATH\n\n"
        "## Canary rollout\n\n"
        "Start every production rollout as a canary:\n\n"
        f"```bash\n{MD_CANARY_CMD}\n```\n\n"
        "Watch the error budget for 20 minutes before promoting.\n\n"
        "## Promotion\n\n"
        "```bash\nlarkctl rollout promote --to 100\n```\n\n"
        "## Rollback\n\n"
        "```yaml\nrollback:\n  strategy: instant\n  keep_canary_logs_days: 14\n```\n\n"
        f"## Notes\n\n{_filler(rng, 160)}\n"
    )
    return CorpusDoc(MD_NAME, "md", "text/markdown", md.encode(), md,
                     _range(_est_chunks(md), 0.5, 3.0),
                     headings=["Prerequisites", "Canary rollout", "Promotion", "Rollback"],
                     heading_probes=[("Canary rollout", "larkctl rollout --canary 10")])


def _font(size: int) -> Any:
    from PIL import ImageFont

    for path in ("/System/Library/Fonts/Supplemental/Arial.ttf",
                 "/System/Library/Fonts/Helvetica.ttc",
                 "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                 "/usr/share/fonts/dejavu/DejaVuSans.ttf"):
        if os.path.exists(path):
            try:
                return ImageFont.truetype(path, size)
            except OSError:
                continue
    try:
        return ImageFont.load_default(size=size)
    except TypeError:  # very old Pillow: fixed bitmap font
        return ImageFont.load_default()


def render_scan(lines: list[str], *, size: tuple[int, int] = (1654, 2339),
                font_px: int = 46, skew: float = 0.6, seed: int = 11) -> Any:
    """Black text on off-white paper with light speckle and a slight skew."""
    from PIL import Image, ImageDraw

    rng = random.Random(seed)
    img = Image.new("L", size, 246)
    draw = ImageDraw.Draw(img)
    font = _font(font_px)
    y = 160
    for line in lines:
        draw.text((140, y), line, fill=20, font=font)
        y += int(font_px * 2.1)
    for _ in range(size[0] * size[1] // 900):  # paper speckle
        draw.point((rng.randrange(size[0]), rng.randrange(size[1])), fill=rng.randint(170, 215))
    return img.rotate(skew, expand=False, fillcolor=246)


def build_scan_pdf() -> CorpusDoc:
    img = render_scan(SCAN_LINES)
    buf = io.BytesIO()
    img.convert("RGB").save(buf, format="PDF", resolution=200.0)
    text = "\n".join(SCAN_LINES)
    return CorpusDoc(SCAN_NAME, "scan_pdf", PDF_MIME, buf.getvalue(), text, (1, 3),
                     pages=1, needs_ocr=True)


def build_png() -> CorpusDoc:
    img = render_scan(PNG_LINES, size=(1400, 900), font_px=54, skew=-0.8, seed=5)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    text = "\n".join(PNG_LINES)
    return CorpusDoc(PNG_NAME, "png", "image/png", buf.getvalue(), text, (1, 2),
                     needs_ocr=True)


def build_zip() -> CorpusDoc:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, content in ZIP_FILES.items():
            z.writestr(name, content)
    text = "\n".join(ZIP_FILES.values())
    return CorpusDoc(ZIP_NAME, "zip", "application/zip", buf.getvalue(), text,
                     (len(ZIP_FILES), 3 * len(ZIP_FILES)))


DEFAULT_SEED = 20261002


def doc_rng(fmt: str, seed: int = DEFAULT_SEED) -> random.Random:
    """Each document's own filler stream, so one document never shifts another's."""
    return random.Random(f"{seed}:{fmt}")


def build_corpus(seed: int = DEFAULT_SEED, pdf_pages: int | None = None) -> list[CorpusDoc]:
    """The whole corpus, deterministic for a given seed and page count."""
    pages = pdf_pages or int(os.getenv("RW_PDF_PAGES", "72"))
    if not 60 <= pages <= 120:
        raise ValueError(f"RW_PDF_PAGES must be 60-120, got {pages}")
    return [build_pdf(doc_rng("pdf", seed), pages), build_docx(doc_rng("docx", seed)),
            build_pptx(doc_rng("pptx", seed)), build_xlsx(doc_rng("xlsx", seed)),
            build_csv(doc_rng("csv", seed)), build_html(doc_rng("html", seed)),
            build_md(doc_rng("md", seed)), build_scan_pdf(), build_png(), build_zip()]


def _number_variants(n: int) -> list[str]:
    """``n`` as raw digits, western grouping and Indian (lakh/crore) grouping."""
    raw = str(n)
    western = f"{n:,}"
    head, tail = raw[:-3], raw[-3:]
    groups: list[str] = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    indian = ",".join([g for g in [head, *groups] if g] + [tail]) if len(raw) > 3 else raw
    return sorted({raw, western, indian})


def load_questions() -> list[dict[str, Any]]:
    """The known-answer questions, with computed answers substituted."""
    with open(os.path.join(FIXTURES, "kb_questions.json"), encoding="utf-8") as fh:
        questions = list(json.load(fh)["questions"])
    for q in questions:
        if "@XLSX_FUEL_TOTAL@" in q.get("answer_any", []):
            q["answer_any"] = _number_variants(xlsx_fuel_total())
    return questions
