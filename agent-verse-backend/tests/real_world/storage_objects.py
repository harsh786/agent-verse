"""Mixed-format object sets for the object-store source scenarios (SRC-OBJ-*).

Every object plants one fact with made-up proper nouns and codes, so a search hit
on it can only come from that object. Builders are deterministic and import
nothing from ``app`` (the scenarios drive the live stack over HTTP only).
"""

from __future__ import annotations

import csv
import io
import json
import zipfile
from dataclasses import dataclass

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class SeedObject:
    key: str  # relative to the scenario prefix
    data: bytes
    content_type: str
    fact: str  # a sentence (or row text) the object carries
    probe: str  # a search query that should rank this object first


def _pdf(lines: list[str]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in lines:
        pdf.multi_cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def _docx(title: str, paragraphs: list[str]) -> bytes:
    from docx import Document

    d = Document()
    d.add_heading(title, level=1)
    for p in paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx(title: str, body: str, notes: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = title
    tb = s.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(8.5), Inches(3))
    tb.text_frame.text = body
    s.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def _xlsx(rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Lane rates"
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv(rows: list[list[str]]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode()


FILLER = (
    "This note is part of the Larkspur depot operations archive. It records routine "
    "observations from the shift and is kept for audit purposes only."
)


def mixed_objects() -> list[SeedObject]:
    """Twelve objects across ten formats, some in nested 'folders'."""
    pdf_fact = ("Berth window code KX-4417: the Tuticorin feeder vessel Marisol Dawn "
                "berths at 04:30 on Thursdays.")
    docx_fact = "The Halvorsen cold-chain clause requires reefer probes logged every 7 minutes."
    pptx_fact = "Project Kestrel Loop sets the Ranchi depot on-time dispatch target at 96.2 percent."
    html_fact = "Gate pass scheme Orchid-9 replaces paper passes at the Hosur yard from week 44."
    md_fact = "To restart the Quillon sorter, hold the amber reset for 12 seconds."
    txt_fact = "The night-shift supervisor for the Bhiwandi annex is Rukmini Desai until December."
    json_fact = "Fuel surcharge band F7 applies above 104 rupees per litre."
    utf8_fact = "Le quai numéro 5 à Lyon ferme à 21h; 横浜倉庫の出荷締切は17時です。"
    nested_fact = "The Kochi transshipment buffer for Pallikkara Spices is 36 hours."
    html = (f"<html><head><title>Hosur yard bulletin</title></head><body><main>"
            f"<h1>Hosur yard bulletin</h1><p>{html_fact}</p><p>{FILLER}</p></main></body></html>")
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("handover/vapi.md", "# Vapi handover\n\nThe Vapi chemical bay key is held "
                    "by gatekeeper Ishaan Rao (badge VB-2291).\n")
    return [
        SeedObject("manuals/berths.pdf", _pdf(["Port berth schedule", pdf_fact, FILLER]),
                   "application/pdf", pdf_fact, "Marisol Dawn berth window Tuticorin"),
        SeedObject("contracts/halvorsen.docx",
                   _docx("Halvorsen cold-chain addendum", [docx_fact, FILLER]), DOCX, docx_fact,
                   "Halvorsen reefer probe logging interval"),
        SeedObject("decks/kestrel.pptx",
                   _pptx("Kestrel Loop review", pptx_fact, "Speaker notes: " + FILLER), PPTX,
                   pptx_fact, "Kestrel Loop Ranchi on-time dispatch target"),
        SeedObject("rates/lanes.xlsx", _xlsx([["Lane", "Carrier", "Rate INR"],
                                              ["Vizag-Raipur", "Tamarind Haulage", "41750"],
                                              ["Surat-Indore", "Copperleaf Roadlines", "38200"]]),
                   XLSX, "Tamarind Haulage",
                   "Tamarind Haulage Vizag-Raipur rate"),
        SeedObject("finance/invoices.csv", _csv([["invoice_id", "customer", "amount_inr"],
                                                 ["INV-88213", "Pellucid Agro", "129400"],
                                                 ["INV-88214", "Saffronline Mills", "8800"]]),
                   "text/csv", "INV-88213", "Pellucid Agro invoice amount"),
        SeedObject("bulletins/hosur.html", html.encode(), "text/html", html_fact,
                   "Orchid-9 gate pass scheme Hosur"),
        SeedObject("runbooks/quillon.md", f"# Quillon sorter\n\n{md_fact}\n\n{FILLER}\n".encode(),
                   "text/markdown", md_fact, "how to restart the Quillon sorter"),
        # Served as a generic type: the format must come from the key / content.
        SeedObject("rosters/bhiwandi.txt", f"{txt_fact}\n{FILLER}\n".encode(),
                   "binary/octet-stream", txt_fact, "Bhiwandi annex night-shift supervisor"),
        SeedObject("policies/fuel.json",
                   json.dumps({"policy": "fuel-surcharge", "rule": json_fact,
                               "owner": "Treasury desk"}).encode(),
                   "application/json", json_fact, "fuel surcharge band F7"),
        SeedObject("intl/lyon-yokohama.md", f"# Quais\n\n{utf8_fact}\n\n{FILLER}\n".encode(),
                   "text/markdown; charset=utf-8", utf8_fact, "横浜倉庫 出荷締切"),
        SeedObject("ops/regions/south/kochi.md", f"# Kochi\n\n{nested_fact}\n\n{FILLER}\n".encode(),
                   "text/markdown", nested_fact, "Kochi transshipment buffer Pallikkara"),
        SeedObject("handover/vapi.zip", zip_buf.getvalue(), "application/zip",
                   "Vapi chemical bay key is held by gatekeeper Ishaan Rao",
                   "who holds the Vapi chemical bay key"),
    ]


def daily_notes(n: int, start: int = 0) -> list[SeedObject]:
    """``n`` small Markdown notes, each with its own fact (bulk 'many objects')."""
    out: list[SeedObject] = []
    for i in range(start, start + n):
        fact = (f"Daily note {i:04d}: dock door D{i % 37 + 1} at the Zirakpur cross-dock was "
                f"inspected by crew Tal-{i:04d}.")
        out.append(SeedObject(f"daily/{i:04d}.md", f"# Note {i:04d}\n\n{fact}\n".encode(),
                              "text/markdown", fact, f"crew Tal-{i:04d} inspection"))
    return out


def big_csv(target_bytes: int, marker_row: int, marker: str) -> bytes:
    """A CSV of roughly ``target_bytes`` with ``marker`` in row ``marker_row``."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["seq", "consignment", "origin", "destination", "remarks"])
    i = 0
    while buf.tell() < target_bytes:
        remarks = marker if i == marker_row else "routine movement, no exceptions recorded"
        w.writerow([i, f"CN-{700000 + i}", "Nhava Sheva", "Ludhiana ICD", remarks])
        i += 1
    return buf.getvalue().encode()
