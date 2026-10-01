"""Realistic KB fixture documents (generated in memory), one distinctive fact each.

Every document carries a made-up proper noun (Halcyon, Zephyrine, ...) so a search
hit can only come from the right document, and a goal answer can only come from
the knowledge base, never from the model's training data.
"""

from __future__ import annotations

import io
from dataclasses import dataclass


@dataclass(frozen=True)
class Doc:
    filename: str
    mime: str
    data: bytes
    query: str  # natural-language search that must hit this doc
    fact: str  # substring the retrieved chunk must contain


PDF_TEXT = (
    "Infrastructure change notice - Project Halcyon\n"
    "The Project Halcyon database migration window is Saturday 14 November 2026, "
    "02:00-05:00 IST. Change owner: Arjun Mehta. Rollback decision point: 04:15 IST."
)
DOCX_TEXT = (
    "Vendor contract summary. The Zephyrine Analytics contract auto-renews on 31 March 2027 "
    "unless it is cancelled in writing at least 60 days before the renewal date. "
    "Annual value: INR 18.4 lakh."
)
PPTX_TITLE = "Quokka Pay - FY27 launch plan"
PPTX_BODY = (
    "Quokka Pay launches first in Coimbatore, targeting 12,500 merchants in Q1 FY27, "
    "followed by Madurai in Q2."
)
XLSX_ROWS = [
    ["warehouse_code", "name", "city", "capacity_pallets"],
    ["BLR-7", "Marigold Fulfilment Centre", "Hoskote", 8400],
    ["HYD-2", "Saffron Hub", "Shamshabad", 6100],
]
CSV_TEXT = (
    "service,owning_team,escalation_contact,pager_hours\n"
    "ledger-sync,Team Ibex,Priya Raman,24x7\n"
    "kyc-gateway,Team Oryx,Daniel Dsouza,business-hours\n"
)
HTML_TEXT = (
    "<html><head><title>Workplace handbook</title></head><body>"
    "<h1>Office network</h1>"
    "<p>The guest Wi-Fi network for the Indiranagar office is called Tamarind-Guest; "
    "its access code rotates every Monday at 09:00.</p></body></html>"
)
MD_TEXT = (
    "# Travel expense policy\n\n"
    "Policy reference **Nightjar-4**. The meal per diem for Pune trips is INR 1,850 per day; "
    "receipts are required above INR 500.\n"
)


def _pdf() -> bytes:
    try:
        from fpdf import FPDF

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.multi_cell(0, 8, PDF_TEXT)
        return bytes(pdf.output())
    except ImportError:  # hand-built single-page PDF
        lines = [ln.replace("(", "[").replace(")", "]") for ln in PDF_TEXT.split("\n")]
        ops = "BT /F1 11 Tf 50 760 Td 14 TL " + " ".join(f"({ln}) '" for ln in lines) + " ET"
        content = ops.encode()
        objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
                b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
                b"/Resources << /Font << /F1 5 0 R >> >> >>",
                b"<< /Length %d >>stream\n" % len(content) + content + b"\nendstream",
                b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
        out = b"%PDF-1.4\n"
        offs = []
        for i, o in enumerate(objs, 1):
            offs.append(len(out))
            out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
        xref = len(out)
        out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
        out += b"".join(b"%010d 00000 n \n" % o for o in offs)
        out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
            len(objs) + 1, xref)
        return out


def _docx() -> bytes:
    from docx import Document

    d = Document()
    d.add_heading("Vendor contracts - renewals", level=1)
    d.add_paragraph(DOCX_TEXT)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = PPTX_TITLE
    tb = s.shapes.add_textbox(Inches(0.7), Inches(2), Inches(8.5), Inches(2))
    tb.text_frame.text = PPTX_BODY
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "warehouses"
    for row in XLSX_ROWS:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_docs() -> list[Doc]:
    return [
        Doc("halcyon-change-notice.pdf", "application/pdf", _pdf(),
            "When is the Project Halcyon database migration window?", "Halcyon"),
        Doc("zephyrine-contract.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document", _docx(),
            "When does the Zephyrine Analytics contract auto-renew?", "Zephyrine"),
        Doc("quokka-pay-launch.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation", _pptx(),
            "Which city does Quokka Pay launch in first?", "Coimbatore"),
        Doc("warehouses.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _xlsx(),
            "What is the pallet capacity of the Marigold Fulfilment Centre?", "Marigold"),
        Doc("service-owners.csv", "text/csv", CSV_TEXT.encode(),
            "Who is the escalation contact for ledger-sync owned by Team Ibex?", "Ibex"),
        Doc("workplace-handbook.html", "text/html", HTML_TEXT.encode(),
            "What is the guest Wi-Fi network called in the Indiranagar office?", "Tamarind"),
        Doc("travel-policy.md", "text/markdown", MD_TEXT.encode(),
            "What is the meal per diem for Pune trips under policy Nightjar-4?", "Nightjar"),
    ]
