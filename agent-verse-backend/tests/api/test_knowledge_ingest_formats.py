"""POST /knowledge/ingest (JSON) and /knowledge/ingest/file format contract.

* Empty / whitespace-only JSON content was a 201 with ``chunks_created: 0`` —
  a fake success. It is now a 422.
* ``.pptx`` was refused with 415; its slide text (and speaker notes) is now
  extracted. Legacy binary ``.ppt`` stays a 415 that says to convert to .pptx.
* Images (``.png/.jpg/.jpeg/.webp``) were refused with 415; they are now OCR'd
  with the existing ``OcrEngine`` (Tesseract, or a vision-capable provider).
  With neither available the upload is an honest 503 — never a stored
  placeholder like "[Image content - vision provider not configured]".
"""

from __future__ import annotations

import io
from typing import Any
from unittest.mock import patch

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.api.knowledge import router as knowledge_router
from app.providers.base import CompletionRequest, CompletionResponse
from app.providers.fake import FakeProvider
from app.rag.semantic_cache import SemanticCache
from app.rag.store import KnowledgeStore
from app.tenancy.context import PlanTier, TenantContext
from app.tenancy.middleware import SecurityHeadersMiddleware, TenantMiddleware

_CTX = TenantContext(tenant_id="tid-fmt", plan=PlanTier.PROFESSIONAL, api_key_id="kid-1")
_KEY = "av_test_fmtkey"
_H = {"X-API-Key": _KEY}


def _app(*, embedder: Any = "fake", llm_provider: Any = None) -> FastAPI:
    app = FastAPI()

    async def _resolve(key: str) -> TenantContext | None:
        return _CTX if key == _KEY else None

    app.add_middleware(TenantMiddleware, key_resolver=_resolve)
    app.add_middleware(SecurityHeadersMiddleware)
    app.include_router(knowledge_router)
    app.state.knowledge_store = KnowledgeStore()
    app.state.semantic_cache = SemanticCache()
    app.state.embedder = FakeProvider(embed_dim=768) if embedder == "fake" else embedder
    app.state.llm_provider = llm_provider
    return app


def _collection(client: TestClient) -> str:
    r = client.post("/knowledge/collections", json={"name": "fmt"}, headers=_H)
    assert r.status_code == 201, r.text
    return str(r.json()["collection_id"])


def _stored_texts(app: FastAPI) -> list[str]:
    store: KnowledgeStore = app.state.knowledge_store
    return [c.content for entry in store._data.values() for c in entry.chunks]


def _upload(
    client: TestClient, cid: str, name: str, data: bytes, mime: str
) -> Any:
    return client.post(
        "/knowledge/ingest/file",
        files={"file": (name, io.BytesIO(data), mime)},
        data={"collection_id": cid},
        headers=_H,
    )


# ── 4. empty JSON content ────────────────────────────────────────────────────


@pytest.mark.parametrize("content", ["", "   ", "\n\t  \n"])
def test_empty_or_whitespace_json_content_is_422(content: str) -> None:
    app = _app()
    client = TestClient(app, raise_server_exceptions=False)
    cid = _collection(client)
    r = client.post(
        "/knowledge/ingest", json={"collection_id": cid, "content": content}, headers=_H
    )
    assert r.status_code == 422, r.text
    assert "empty" in r.json()["detail"].lower()
    assert _stored_texts(app) == []


def test_non_empty_json_content_still_ingests() -> None:
    app = _app()
    client = TestClient(app, raise_server_exceptions=False)
    cid = _collection(client)
    r = client.post(
        "/knowledge/ingest",
        json={"collection_id": cid, "content": "The VPN rotates keys every 90 days."},
        headers=_H,
    )
    assert r.status_code == 201, r.text
    assert r.json()["chunks_created"] == 1


# ── 5a. PPTX ─────────────────────────────────────────────────────────────────

_PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Q3 Revenue Review"
    s1.placeholders[1].text = "APAC revenue grew 18% to INR 42 crore."
    s1.notes_slide.notes_text_frame.text = "Mention the Pune office launch."
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Headcount"
    table = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for r, row in enumerate([["Team", "Engineers"], ["Payments", "37"]]):
        for c, v in enumerate(row):
            table.cell(r, c).text = v
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_pptx_slide_text_notes_and_tables_are_ingested() -> None:
    app = _app()
    client = TestClient(app, raise_server_exceptions=False)
    cid = _collection(client)
    r = _upload(client, cid, "review.pptx", _pptx(), _PPTX_MIME)
    assert r.status_code == 201, r.text
    assert r.json()["chunks_created"] >= 1
    text = "\n".join(_stored_texts(app))
    for needle in ("Q3 Revenue Review", "INR 42 crore", "Pune office launch", "Payments", "37"):
        assert needle in text, text


def test_corrupt_pptx_is_422() -> None:
    client = TestClient(_app(), raise_server_exceptions=False)
    cid = _collection(client)
    r = _upload(client, cid, "broken.pptx", b"PK\x03\x04 not a deck", _PPTX_MIME)
    assert r.status_code == 422, r.text


def test_legacy_ppt_is_415_asking_for_pptx() -> None:
    client = TestClient(_app(), raise_server_exceptions=False)
    cid = _collection(client)
    r = _upload(client, cid, "old.ppt", b"\xd0\xcf\x11\xe0 legacy", "application/vnd.ms-powerpoint")
    assert r.status_code == 415, r.text
    assert ".pptx" in r.json()["detail"]


# ── 5b. Images via OCR ───────────────────────────────────────────────────────


def _png(text: str = "INVOICE 4471") -> bytes:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (320, 80), "white")
    ImageDraw.Draw(img).text((10, 30), text, fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _jpeg() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (40, 40), "white").save(buf, format="JPEG")
    return buf.getvalue()


class _VisionProvider:
    def __init__(self, text: str) -> None:
        self._text = text
        self.calls = 0

    def supports_vision(self) -> bool:
        return True

    async def complete(self, request: CompletionRequest) -> CompletionResponse:
        self.calls += 1
        assert request.messages[0].image_data
        return CompletionResponse(content=self._text, model="vision-test")


def test_image_without_ocr_or_vision_is_503_and_stores_nothing() -> None:
    app = _app(llm_provider=None)
    client = TestClient(app, raise_server_exceptions=False)
    cid = _collection(client)
    with patch("app.ingestion.document_text.tesseract_available", return_value=False):
        r = _upload(client, cid, "scan.png", _png(), "image/png")
    assert r.status_code == 503, r.text
    detail = r.json()["detail"]
    assert "OCR" in detail and "vision" in detail
    assert _stored_texts(app) == []


def test_fake_provider_is_not_treated_as_a_vision_model() -> None:
    app = _app(llm_provider=FakeProvider())
    client = TestClient(app, raise_server_exceptions=False)
    cid = _collection(client)
    with patch("app.ingestion.document_text.tesseract_available", return_value=False):
        r = _upload(client, cid, "scan.png", _png(), "image/png")
    assert r.status_code == 503, r.text
    assert _stored_texts(app) == []


def test_image_is_ocrd_with_tesseract_when_available() -> None:
    app = _app()
    client = TestClient(app, raise_server_exceptions=False)
    cid = _collection(client)

    async def _page(self: Any, img: Any, *, provider: Any = None, vision_fallback: bool = True):
        return "INVOICE 4471 total INR 4,200 due 30 Nov", 0.93, "tesseract"

    with (
        patch("app.ingestion.document_text.tesseract_available", return_value=True),
        patch("app.ocr.engine.OcrEngine._ocr_page", _page),
    ):
        r = _upload(client, cid, "invoice.png", _png(), "image/png")
    assert r.status_code == 201, r.text
    assert r.json()["chunks_created"] == 1
    assert _stored_texts(app) == ["INVOICE 4471 total INR 4,200 due 30 Nov"]
    store: KnowledgeStore = app.state.knowledge_store
    chunk = next(iter(store._data.values())).chunks[0]
    assert chunk.metadata["ocr_engine"] == "tesseract"
    assert chunk.metadata["source_type"] == "image"


@pytest.mark.parametrize(
    ("name", "mime", "data_fn"),
    [("photo.jpg", "image/jpeg", _jpeg), ("photo.jpeg", "image/jpeg", _jpeg)],
)
def test_image_is_ocrd_with_a_vision_provider_without_tesseract(
    name: str, mime: str, data_fn: Any
) -> None:
    vision = _VisionProvider("Whiteboard: launch on 12 December")
    app = _app(llm_provider=vision)
    client = TestClient(app, raise_server_exceptions=False)
    cid = _collection(client)
    with patch("app.ingestion.document_text.tesseract_available", return_value=False):
        r = _upload(client, cid, name, data_fn(), mime)
    assert r.status_code == 201, r.text
    assert vision.calls == 1
    assert _stored_texts(app) == ["Whiteboard: launch on 12 December"]


def test_webp_is_accepted() -> None:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (40, 40), "white").save(buf, format="WEBP")
    vision = _VisionProvider("webp text")
    app = _app(llm_provider=vision)
    client = TestClient(app, raise_server_exceptions=False)
    cid = _collection(client)
    with patch("app.ingestion.document_text.tesseract_available", return_value=False):
        r = _upload(client, cid, "shot.webp", buf.getvalue(), "image/webp")
    assert r.status_code == 201, r.text


def test_image_with_no_extractable_text_is_422_not_a_placeholder() -> None:
    app = _app(llm_provider=_VisionProvider("   "))
    client = TestClient(app, raise_server_exceptions=False)
    cid = _collection(client)
    with patch("app.ingestion.document_text.tesseract_available", return_value=False):
        r = _upload(client, cid, "blank.png", _png(), "image/png")
    assert r.status_code == 422, r.text
    assert _stored_texts(app) == []


def test_corrupt_image_is_422() -> None:
    client = TestClient(_app(), raise_server_exceptions=False)
    cid = _collection(client)
    with patch("app.ingestion.document_text.tesseract_available", return_value=True):
        r = _upload(client, cid, "bad.png", b"\x89PNG\r\n\x1a\n\x00\x00garbage", "image/png")
    assert r.status_code == 422, r.text


def test_image_type_is_inferred_from_mime_when_the_name_has_no_extension() -> None:
    vision = _VisionProvider("pasted screenshot text")
    app = _app(llm_provider=vision)
    client = TestClient(app, raise_server_exceptions=False)
    cid = _collection(client)
    with patch("app.ingestion.document_text.tesseract_available", return_value=False):
        r = _upload(client, cid, "clipboard", _png(), "image/png")
    assert r.status_code == 201, r.text
    assert _stored_texts(app) == ["pasted screenshot text"]


# ── 2. embedder unavailability is explained ─────────────────────────────────


def test_ingest_503_explains_why_the_embedder_is_unavailable() -> None:
    from app.providers.embedder_factory import EmbedderResolution

    app = _app(embedder=None)
    app.state.embedder_resolution = EmbedderResolution(
        embedder=None,
        errors=[("sentence_transformers", "OSError: all-mpnet-base-v3 is not a valid model")],
    )
    client = TestClient(app, raise_server_exceptions=False)
    cid = _collection(client)
    r = client.post(
        "/knowledge/ingest", json={"collection_id": cid, "content": "hello world"}, headers=_H
    )
    assert r.status_code == 503
    detail = r.json()["detail"]
    assert detail.startswith("Embedding provider is unavailable")
    assert "sentence_transformers" in detail and "not a valid model" in detail


def test_ingest_503_says_not_configured_when_nothing_is_set() -> None:
    from app.providers.embedder_factory import EmbedderResolution

    app = _app(embedder=None)
    app.state.embedder_resolution = EmbedderResolution(embedder=None)
    client = TestClient(app, raise_server_exceptions=False)
    cid = _collection(client)
    r = client.post(
        "/knowledge/ingest", json={"collection_id": cid, "content": "hello world"}, headers=_H
    )
    assert r.status_code == 503
    assert "not configured" in r.json()["detail"]
