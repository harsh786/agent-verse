"""P1d (A10): ``POST /knowledge/ingest/url`` and document re-ingestion.

Before P1d, URL ingestion:

* decoded every response as text — a PDF or DOCX served by URL was indexed as
  binary garbage (P1d-2);
* used the public-only SSRF guard, so an operator-allowlisted internal host
  (the ingestion allowlist every connector honours) could not be ingested by
  URL, and a permanent move was not reported (P1d-1);
* answered 500 "Failed to fetch" for every upstream failure — 404, 500,
  timeouts and refused redirects alike (P1d-3);
* decoded HTML as UTF-8 regardless of its charset declaration (P1d-4);
* replaced a document under legal hold without checking the hold (P1d-5).

The real fetch path runs here (``httpx.AsyncClient.send`` is the only fake), so
the egress guard, redirects, the size cap and charset handling are exercised.
"""

from __future__ import annotations

import io
import socket
from collections.abc import Iterator
from typing import Any
from unittest.mock import AsyncMock, patch

import httpx
import pytest
from fastapi.testclient import TestClient

from app.api.knowledge import stable_url_document_id
from app.rag.store import KnowledgeStore
from tests.api.test_knowledge_extra4 import _CTX, H, _create_collection, _make_app

PAGE = "https://docs.example.org/handbook/depot-hours"

_Route = tuple[int, dict[str, str], bytes]


@pytest.fixture(autouse=True)
def _public_dns(monkeypatch: pytest.MonkeyPatch) -> None:
    real = socket.getaddrinfo

    def fake(host: Any, *args: Any, **kwargs: Any) -> Any:
        if str(host).endswith("example.org"):
            return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 0))]
        return real(host, *args, **kwargs)

    monkeypatch.setattr(socket, "getaddrinfo", fake)


@pytest.fixture
def site() -> Iterator[dict[str, _Route]]:
    routes: dict[str, _Route] = {}

    async def _send(self: httpx.AsyncClient, request: httpx.Request, **_: Any) -> httpx.Response:
        url = str(request.url)
        if url not in routes:
            return httpx.Response(404, content=b"not here", request=request)
        status, headers, body = routes[url]
        if body == b"<timeout>":
            raise httpx.ReadTimeout("read timed out", request=request)
        return httpx.Response(status, headers=headers, content=body, request=request)

    with patch.object(httpx.AsyncClient, "send", _send):
        yield routes


def _client(store: KnowledgeStore | None = None) -> tuple[TestClient, KnowledgeStore, str]:
    from app.providers.fake import FakeProvider

    store = store or KnowledgeStore()
    client = TestClient(_make_app(knowledge_store=store, embedder=FakeProvider(embed_dim=8)),
                        raise_server_exceptions=False)
    return client, store, _create_collection(client)


def _ingest(client: TestClient, cid: str, url: str = PAGE) -> httpx.Response:
    return client.post("/knowledge/ingest/url", headers=H,
                       json={"collection_id": cid, "url": url, "source_type": "web"})


def _texts(store: KnowledgeStore, cid: str) -> str:
    return "\n".join(c.content for c in store._data[(_CTX.tenant_id, cid)].chunks)


def _html(body: str, head: str = "") -> bytes:
    return (f"<!doctype html><html><head>{head}<title>Depot hours</title>"
            "<style>.x{color:red}</style><script>window.dataLayer=[];</script></head><body>"
            "<nav><a href='/'>Home</a> <a href='/careers'>Careers at Kestrel</a></nav>"
            "<div id='cookie-banner'>We use cookies to improve your experience</div>"
            f"<main><h1>Depot hours</h1>{body}</main>"
            "<footer>Copyright Kestrel Logistics. Subscribe to our newsletter</footer>"
            "</body></html>").encode()


# ── P1d-2: documents served by URL ──────────────────────────────────────────


def _pdf(lines: list[list[str]]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    for page in lines:
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        for line in page:
            pdf.cell(0, 10, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def _docx(paragraphs: list[str]) -> bytes:
    import docx

    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_a_pdf_served_by_url_is_read_page_by_page(site: dict[str, _Route]) -> None:
    url = "https://docs.example.org/download?id=tariff-2026"
    site[url] = (200, {"content-type": "application/octet-stream"}, _pdf([
        ["Tariff schedule 2026", "Reefer plug-in per day: 2,450 INR"],
        ["Penalties", "Late gate-in after cut-off: 9,800 INR per box"],
    ]))
    client, store, cid = _client()
    resp = _ingest(client, cid, url)
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["content_kind"] == "pdf"
    assert body["pages"] == 2
    chunks = store._data[(_CTX.tenant_id, cid)].chunks
    late = next(c for c in chunks if "Late gate-in after cut-off" in c.content)
    assert late.metadata["page"] == "2"
    assert late.metadata["source_url"] == url
    assert "%PDF" not in _texts(store, cid)


def test_a_docx_served_by_url_is_extracted(site: dict[str, _Route]) -> None:
    url = "https://docs.example.org/policies/handover.docx"
    site[url] = (200, {"content-type": "application/vnd.openxmlformats-officedocument."
                                        "wordprocessingml.document"},
                 _docx(["Shift handover policy", "Seal numbers are re-checked by two people."]))
    client, store, cid = _client()
    resp = _ingest(client, cid, url)
    assert resp.status_code == 201, resp.text
    assert resp.json()["content_kind"] == "docx"
    text = _texts(store, cid)
    assert "Seal numbers are re-checked by two people." in text
    assert "word/document.xml" not in text and "PK" not in text[:4]


def test_plain_text_by_url_is_kept_verbatim(site: dict[str, _Route]) -> None:
    url = "https://docs.example.org/notes/changelog.txt"
    site[url] = (200, {"content-type": "text/plain; charset=utf-8"},
                 b"Version 4.2: gate scanners report seal numbers.\n")
    client, store, cid = _client()
    assert _ingest(client, cid, url).status_code == 201
    assert "Version 4.2: gate scanners report seal numbers." in _texts(store, cid)


def test_an_unsupported_binary_is_refused_not_indexed(site: dict[str, _Route]) -> None:
    url = "https://docs.example.org/files/blob.bin"
    site[url] = (200, {"content-type": "application/octet-stream"}, b"\x00\x01\x02" * 400)
    client, store, cid = _client()
    resp = _ingest(client, cid, url)
    assert resp.status_code == 415, resp.text
    assert store._data[(_CTX.tenant_id, cid)].chunks == []


# ── boilerplate and charset (P1d-4) ─────────────────────────────────────────


def test_html_boilerplate_is_dropped(site: dict[str, _Route]) -> None:
    site[PAGE] = (200, {"content-type": "text/html"},
                  _html("<p>The Hosur depot opens at 06:30 on weekdays.</p>"))
    client, store, cid = _client()
    resp = _ingest(client, cid)
    assert resp.status_code == 201, resp.text
    assert resp.json()["title"] == "Depot hours"
    text = _texts(store, cid)
    assert "The Hosur depot opens at 06:30 on weekdays." in text
    for chrome in ("Careers at Kestrel", "We use cookies", "Subscribe to our newsletter",
                   "dataLayer", "color:red"):
        assert chrome not in text


@pytest.mark.parametrize(
    ("encoding", "header", "meta", "sentence"),
    [
        ("cp1252", "text/html", '<meta charset="windows-1252">',
         "Lieferung für Zürich: 12 € – Smörgåsbord café"),
        ("shift_jis", "text/html; charset=Shift_JIS", "", "東京都江東区の倉庫は午前九時に開きます"),
        ("koi8_r", "text/html",
         '<meta http-equiv="Content-Type" content="text/html; charset=KOI8-R">',
         "Склад в Новосибирске открыт с девяти утра"),
        ("utf-8", "text/html", "", "दिल्ली गोदाम सुबह नौ बजे खुलता है"),
    ],
)
def test_html_charset_variants(site: dict[str, _Route], encoding: str, header: str, meta: str,
                               sentence: str) -> None:
    site[PAGE] = (200, {"content-type": header},
                  _html(f"<p>{sentence}</p>", head=meta).decode().encode(encoding))
    client, store, cid = _client()
    assert _ingest(client, cid).status_code == 201
    assert sentence in _texts(store, cid)


# ── P1d-1 / P1d-3: redirects and honest failures ────────────────────────────


def test_permanent_redirect_is_followed_and_reported(site: dict[str, _Route]) -> None:
    new = "https://docs.example.org/handbook/v2/depot-hours"
    site[PAGE] = (301, {"location": "/handbook/v2/depot-hours"}, b"")
    site[new] = (200, {"content-type": "text/html"}, _html("<p>Moved page body text.</p>"))
    client, store, cid = _client()
    resp = _ingest(client, cid)
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["final_url"] == new
    assert body["moved_permanently"] == {"from": PAGE, "to": new, "status": 301}
    # The document stays the requested URL's (stable id); the move is on its chunks.
    assert body["document_id"] == stable_url_document_id(_CTX.tenant_id, cid, PAGE)
    chunk = store._data[(_CTX.tenant_id, cid)].chunks[0]
    assert chunk.metadata["moved_to"] == new


def test_redirect_to_an_internal_address_is_refused(site: dict[str, _Route]) -> None:
    site[PAGE] = (302, {"location": "http://169.254.169.254/latest/meta-data/"}, b"")
    client, store, cid = _client()
    resp = _ingest(client, cid)
    assert resp.status_code == 400
    assert "blocked" in resp.json()["detail"]
    assert store._data[(_CTX.tenant_id, cid)].chunks == []


@pytest.mark.parametrize(("status", "expected"), [(404, 502), (410, 502), (500, 502),
                                                  (503, 502)])
def test_upstream_error_status_is_reported_not_500(site: dict[str, _Route], status: int,
                                                   expected: int) -> None:
    site[PAGE] = (status, {}, b"upstream says no")
    client, _store, cid = _client()
    resp = _ingest(client, cid)
    assert resp.status_code == expected
    assert f"HTTP {status}" in resp.json()["detail"]


def test_timeout_is_504(site: dict[str, _Route]) -> None:
    site[PAGE] = (200, {}, b"<timeout>")
    client, _store, cid = _client()
    resp = _ingest(client, cid)
    assert resp.status_code == 504
    assert "timed out" in resp.json()["detail"]


def test_redirect_loop_is_422(site: dict[str, _Route]) -> None:
    other = "https://docs.example.org/loop"
    site[PAGE] = (302, {"location": other}, b"")
    site[other] = (302, {"location": PAGE}, b"")
    client, _store, cid = _client()
    resp = _ingest(client, cid)
    assert resp.status_code == 422
    assert "redirect" in resp.json()["detail"]


def test_a_body_over_the_limit_is_413(site: dict[str, _Route],
                                      monkeypatch: pytest.MonkeyPatch) -> None:
    from app.core.config import get_settings

    monkeypatch.setattr(get_settings(), "knowledge_max_upload_bytes", 2048)
    site[PAGE] = (200, {"content-type": "text/html"}, _html("<p>" + "x" * 5000 + "</p>"))
    client, store, cid = _client()
    resp = _ingest(client, cid)
    assert resp.status_code == 413
    assert store._data[(_CTX.tenant_id, cid)].chunks == []


def test_an_operator_allowlisted_internal_host_is_ingested(
    site: dict[str, _Route], monkeypatch: pytest.MonkeyPatch
) -> None:
    """URL ingestion honours the ingestion egress allowlist (both env switches),
    like every connector — and only for the listed host."""
    from app.core.config import get_settings

    settings = get_settings()
    monkeypatch.setattr(settings, "ingestion_allow_internal_sources", True)
    monkeypatch.setattr(settings, "ingestion_internal_source_allowlist", "wiki.corp.internal")
    real = socket.getaddrinfo

    def fake(host: Any, *args: Any, **kwargs: Any) -> Any:
        if str(host) in ("wiki.corp.internal", "other.corp.internal"):
            return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("10.20.0.7", 0))]
        return real(host, *args, **kwargs)

    monkeypatch.setattr(socket, "getaddrinfo", fake)
    allowed = "http://wiki.corp.internal/runbooks/gate"
    site[allowed] = (200, {"content-type": "text/plain"}, b"Gate runbook: rotate the seals.")
    client, store, cid = _client()
    assert _ingest(client, cid, allowed).status_code == 201
    blocked = _ingest(client, cid, "http://other.corp.internal/x")
    assert blocked.status_code == 400


# ── P1d-5: re-ingest replaces; legal hold blocks it ─────────────────────────


def _hold(client: TestClient, held_ids: set[str]) -> None:
    async def is_under_hold(*, resource_id: str, tenant_id: str) -> bool:
        return resource_id in held_ids

    mgr = AsyncMock()
    mgr.is_under_hold = is_under_hold
    client.app.state.legal_hold_manager = mgr  # type: ignore[attr-defined]


def test_reingest_replaces_the_document(site: dict[str, _Route]) -> None:
    site[PAGE] = (200, {"content-type": "text/html"}, _html("<p>Opens at 06:30.</p>"))
    client, store, cid = _client()
    first = _ingest(client, cid).json()
    site[PAGE] = (200, {"content-type": "text/html"}, _html("<p>Opens at 07:15 from May.</p>"))
    second = _ingest(client, cid)
    assert second.status_code == 201
    body = second.json()
    assert body["document_id"] == first["document_id"]
    assert body["replaced"] is True
    text = _texts(store, cid)
    assert "07:15" in text and "06:30" not in text
    unchanged = _ingest(client, cid).json()
    assert unchanged["deduplicated"] is True
    assert unchanged["document_id"] == first["document_id"]


def test_reingest_of_a_held_document_is_409_and_keeps_it(site: dict[str, _Route]) -> None:
    site[PAGE] = (200, {"content-type": "text/html"}, _html("<p>Opens at 06:30.</p>"))
    client, store, cid = _client()
    doc_id = _ingest(client, cid).json()["document_id"]
    _hold(client, {doc_id})
    site[PAGE] = (200, {"content-type": "text/html"}, _html("<p>Opens at 07:15 from May.</p>"))
    resp = _ingest(client, cid)
    assert resp.status_code == 409
    assert "legal hold" in resp.json()["detail"]
    text = _texts(store, cid)
    assert "06:30" in text and "07:15" not in text
    again = client.post(f"/knowledge/collections/{cid}/documents/{doc_id}/reingest", headers=H)
    assert again.status_code == 409
    assert "06:30" in _texts(store, cid)


def test_a_held_collection_blocks_replacing_its_url_documents(site: dict[str, _Route]) -> None:
    site[PAGE] = (200, {"content-type": "text/html"}, _html("<p>Opens at 06:30.</p>"))
    client, store, cid = _client()
    assert _ingest(client, cid).status_code == 201
    _hold(client, {cid})
    site[PAGE] = (200, {"content-type": "text/html"}, _html("<p>Opens at 07:15.</p>"))
    assert _ingest(client, cid).status_code == 409
    assert "06:30" in _texts(store, cid)


def test_a_first_ingest_is_not_blocked_by_an_unrelated_hold(site: dict[str, _Route]) -> None:
    site[PAGE] = (200, {"content-type": "text/html"}, _html("<p>Opens at 06:30.</p>"))
    client, _store, cid = _client()
    _hold(client, {"some-other-document"})
    assert _ingest(client, cid).status_code == 201
