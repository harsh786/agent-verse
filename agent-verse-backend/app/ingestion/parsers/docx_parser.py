"""DOCXParser — extracts structured text from Word documents."""
from __future__ import annotations
import io
from dataclasses import dataclass, field
from typing import Any


@dataclass
class DOCXParseResult:
    source_name: str
    paragraphs: list[str] = field(default_factory=list)
    headings: list[str] = field(default_factory=list)
    error: str | None = None

    def to_chunks(self) -> list[dict[str, Any]]:
        if not self.paragraphs:
            return []
        chunks: list[dict[str, Any]] = []
        current: list[str] = []
        current_len = 0
        for para in self.paragraphs:
            if current_len + len(para) > 2000 and current:
                chunks.append({
                    "content": "\n\n".join(current),
                    "chunk_index": len(chunks),
                    "source_name": self.source_name,
                    "content_type": "docx",
                })
                current = [para]
                current_len = len(para)
            else:
                current.append(para)
                current_len += len(para)
        if current:
            chunks.append({
                "content": "\n\n".join(current),
                "chunk_index": len(chunks),
                "source_name": self.source_name,
                "content_type": "docx",
            })
        return chunks


class DOCXParser:
    def parse_bytes(self, docx_bytes: bytes, source_name: str) -> DOCXParseResult:
        if not docx_bytes:
            return DOCXParseResult(source_name=source_name, error="empty document")
        try:
            from docx import Document  # type: ignore[import]
            doc = Document(io.BytesIO(docx_bytes))
            paragraphs: list[str] = []
            headings: list[str] = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                if para.style.name.startswith("Heading"):
                    headings.append(text)
                paragraphs.append(text)
            return DOCXParseResult(source_name=source_name, paragraphs=paragraphs, headings=headings)
        except ImportError:
            text = docx_bytes.decode("utf-8", errors="replace")
            paras = [p.strip() for p in text.split("\n\n") if p.strip()]
            return DOCXParseResult(source_name=source_name, paragraphs=paras)
        except Exception as exc:
            return DOCXParseResult(source_name=source_name, error=str(exc))
