"""Text extraction for binary documents (PDF, DOCX, PPTX, images) — one fail-closed path.

Every upload path used to "degrade" when a parser was missing or failed: the
raw bytes were decoded as UTF-8 (so a PDF was indexed as ``%PDF-1.3 … endobj``)
or a placeholder chunk ("install pypdf") was embedded, and the upload reported
success. A knowledge base full of PDF syntax retrieves nothing useful and looks
healthy. These helpers either return the document's real text or raise.
"""

from __future__ import annotations

import io
import os
from typing import Any


class DocumentParseError(ValueError):
    """The document could not be parsed, or contains no extractable text."""


class ParserUnavailableError(RuntimeError):
    """The parser library for this format is not installed on this host."""


def extract_pdf_pages(
    data: bytes, *, filename: str = "document.pdf", allow_textless: bool = False
) -> list[str]:
    """Return the text of each page (``""`` for pages without a text layer).

    Raises DocumentParseError for a corrupt/encrypted PDF or one with no text
    at all (e.g. a scan, which needs OCR), ParserUnavailableError without pypdf.
    ``allow_textless=True`` returns the (all empty) pages of a scan instead, for
    a caller that OCRs them (:func:`ocr_pdf_pages`).
    """
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as exc:  # pragma: no cover - pypdf is a core dependency
        raise ParserUnavailableError("PDF parsing requires pypdf") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not _decrypt_without_password(reader):
            raise DocumentParseError(
                f"{filename}: the PDF is encrypted and needs a password to open; upload an "
                "unprotected copy"
            )
        pages = [page.extract_text() or "" for page in reader.pages]
    except DocumentParseError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError) as exc:
        raise DocumentParseError(f"{filename}: not a readable PDF ({exc})") from exc
    if not allow_textless and not any(p.strip() for p in pages):
        raise DocumentParseError(
            f"{filename}: the PDF has no extractable text (scanned images need OCR)"
        )
    return pages


def _decrypt_without_password(reader: Any) -> bool:
    """Open a permissions-only encrypted PDF (empty user password), as every viewer
    does; False when a real password is needed."""
    from pypdf import PasswordType

    try:
        return bool(reader.decrypt("") != PasswordType.NOT_DECRYPTED)
    except Exception as exc:  # e.g. an AES PDF without the cryptography package
        raise DocumentParseError(f"the PDF's encryption could not be opened ({exc})") from exc


# Pages of one upload OCR'd in the request (each is ~1-3 s of Tesseract at 300 dpi).
OCR_MAX_PDF_PAGES = int(os.getenv("KNOWLEDGE_OCR_MAX_PAGES", "60"))
OCR_PDF_DPI = 300


def render_pdf_page(data: bytes, page_number: int, *, dpi: int = OCR_PDF_DPI) -> Any:
    """Page ``page_number`` (1-based) of a PDF as a grayscale PIL image (poppler)."""
    try:
        from pdf2image import convert_from_bytes
        from pdf2image.exceptions import PDFInfoNotInstalledError
    except ImportError as exc:
        raise ParserUnavailableError(
            "scanned-PDF OCR needs pdf2image and poppler (the 'ocr' extra)"
        ) from exc
    try:
        images = convert_from_bytes(
            data, dpi=dpi, first_page=page_number, last_page=page_number, grayscale=True
        )
    except PDFInfoNotInstalledError as exc:
        raise ParserUnavailableError(
            "scanned-PDF OCR needs poppler (pdftoppm), which is not installed"
        ) from exc
    except Exception as exc:
        raise DocumentParseError(
            f"page {page_number} could not be rendered for OCR ({exc})"
        ) from exc
    if not images:
        raise DocumentParseError(f"page {page_number} could not be rendered for OCR")
    return images[0]


def _require_ocr(filename: str, vision_provider: Any) -> bool:
    """True when the vision provider will be used; raise when no OCR is possible."""
    has_vision = is_vision_provider(vision_provider)
    if not has_vision and not tesseract_available():
        raise OcrUnavailableError(
            f"{filename}: this file needs OCR, but no OCR engine is available: install "
            "Tesseract (the tesseract binary plus the 'ocr' extra) or configure a "
            "vision-capable model provider"
        )
    return has_vision


async def ocr_pdf_pages(
    data: bytes,
    *,
    filename: str,
    page_numbers: list[int],
    vision_provider: Any = None,
    ocr_engine: Any = None,
) -> dict[int, tuple[str, str]]:
    """OCR the given 1-based pages of a PDF: ``{page: (text, engine_used)}``.

    Each page is rendered on its own (so one page's text never lands on
    another page's citation). Raises OcrUnavailableError without an OCR engine,
    DocumentParseError when more pages need OCR than one request may run.
    """
    import asyncio

    if len(page_numbers) > OCR_MAX_PDF_PAGES:
        raise DocumentParseError(
            f"{filename}: {len(page_numbers)} pages have no text layer and need OCR, but at "
            f"most {OCR_MAX_PDF_PAGES} scanned pages are OCR'd per upload; split the PDF"
        )
    has_vision = _require_ocr(filename, vision_provider)
    if ocr_engine is None:
        from app.ocr.engine import OcrEngine

        ocr_engine = OcrEngine()
    out: dict[int, tuple[str, str]] = {}
    for number in page_numbers:
        image = await asyncio.to_thread(render_pdf_page, data, number)
        buf = io.BytesIO()
        image.save(buf, format="PNG")
        result = await ocr_engine.extract(
            image_bytes=buf.getvalue(),
            provider=vision_provider if has_vision else None,
            extract_fields=False,
            vision_fallback=has_vision,
        )
        text = (getattr(result, "raw_text", "") or "").strip()
        out[number] = (text, str(getattr(result, "engine_used", "") or ""))
    return out


def extract_docx_text(data: bytes, *, filename: str = "document.docx") -> str:
    """Return the paragraph text of a .docx; raise instead of guessing."""
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - python-docx is a core dependency
        raise ParserUnavailableError("DOCX parsing requires python-docx") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types
        raise DocumentParseError(f"{filename}: not a readable .docx ({exc})") from exc
    text = "\n".join(_docx_blocks(document))
    if not text.strip():
        raise DocumentParseError(f"{filename}: the document has no text")
    return text


def _docx_blocks(document: object) -> list[str]:
    """Paragraphs AND tables in document order.

    ``document.paragraphs`` skips tables entirely, so the figures a policy or a
    price list keeps in a table never reached the index. Each table row is
    emitted as ``header: value`` pairs so it stays searchable after chunking.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body  # type: ignore[attr-defined]
    out: list[str] = []
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)  # type: ignore[arg-type]
            text = paragraph.text
            if text.strip():
                level = _docx_heading_level(paragraph)
                out.append(f"{'#' * level} {text.strip()}" if level else text)
        elif tag == "tbl":
            table = Table(child, document)  # type: ignore[arg-type]
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            header, *data_rows = rows
            if not data_rows:
                out.append(" | ".join(header))
                continue
            for row in data_rows:
                pairs = [
                    f"{header[i] if i < len(header) and header[i] else f'col{i + 1}'}: {v}"
                    for i, v in enumerate(row)
                    if v
                ]
                if pairs:
                    out.append("; ".join(pairs))
    return out


def _docx_heading_level(paragraph: Any) -> int:
    """Markdown level of a Word heading paragraph (Title = 1, Heading N = N), else 0,
    so the upload chunker can start a chunk per section."""
    try:
        name = str(paragraph.style.name or "")
    except Exception:
        return 0
    if name == "Title":
        return 1
    if name.startswith("Heading "):
        digits = name.removeprefix("Heading ").strip()
        if digits.isdigit():
            return max(1, min(6, int(digits)))
    return 0


class UnsupportedDocumentError(ValueError):
    """The file type has no text extractor here (a 415 for an upload)."""


# Known binary office/e-book formats with no extractor installed.
_UNSUPPORTED_BINARY_EXTS = frozenset(
    {"doc", "xls", "odt", "ods", "odp", "rtf", "epub", "pages", "numbers", "key"}
)

# Images accepted by the upload path; their text comes from OCR (async, see
# extract_image_text), never from extract_upload_text.
IMAGE_UPLOAD_EXTS = frozenset({"png", "jpg", "jpeg", "webp"})

# MIME → extension, for uploads whose filename carries no extension.
UPLOAD_MIME_EXTS: dict[str, str] = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.ms-powerpoint": "ppt",
    "image/png": "png",
    "image/jpeg": "jpg",
    "image/jpg": "jpg",
    "image/webp": "webp",
}


class OcrUnavailableError(ParserUnavailableError):
    """No OCR engine (Tesseract) and no vision-capable provider is configured."""


def extract_pptx_text(data: bytes, *, filename: str = "slides.pptx") -> str:
    """Slide text (titles, text boxes, tables, grouped shapes) and speaker notes."""
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - python-pptx is a core dependency
        raise ParserUnavailableError("PPTX parsing requires python-pptx") from exc
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:  # python-pptx raises several unrelated types
        raise DocumentParseError(f"{filename}: not a readable .pptx ({exc})") from exc
    blocks: list[str] = []
    for number, slide in enumerate(prs.slides, start=1):
        lines = [line for shape in slide.shapes for line in _pptx_shape_lines(shape)]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                lines.append(f"Notes: {notes.text.strip()}")
        if lines:
            blocks.append(f"Slide {number}:\n" + "\n".join(lines))
    text = "\n\n".join(blocks)
    if not text.strip():
        raise DocumentParseError(f"{filename}: the presentation has no text")
    return text


def _pptx_shape_lines(shape: Any) -> list[str]:
    from pptx.shapes.group import GroupShape

    lines: list[str] = []
    if isinstance(shape, GroupShape):
        for child in shape.shapes:
            lines.extend(_pptx_shape_lines(child))
        return lines
    if getattr(shape, "has_table", False):
        rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
        lines.extend(" | ".join(c for c in row if c) for row in rows if any(row))
        return lines
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()
        if text:
            lines.append(text)
    return lines


def tesseract_available() -> bool:
    """True when pytesseract AND the tesseract binary are usable on this host."""
    try:
        import pytesseract

        pytesseract.get_tesseract_version()
    except Exception:
        return False
    return True


def is_vision_provider(provider: Any) -> bool:
    """A real provider that says it accepts images (never the canned FakeProvider)."""
    if provider is None:
        return False
    from app.providers.fake import FakeProvider

    if isinstance(provider, FakeProvider):
        return False
    supports = getattr(provider, "supports_vision", None)
    try:
        return bool(callable(supports) and supports())
    except Exception:
        return False


async def extract_image_text(
    data: bytes,
    *,
    filename: str,
    vision_provider: Any = None,
    ocr_engine: Any = None,
) -> tuple[str, str]:
    """OCR an image with the existing OcrEngine; return ``(text, engine_used)``.

    Raises DocumentParseError (unreadable image, or no text found) or
    OcrUnavailableError when neither Tesseract nor a vision-capable provider is
    available — never a placeholder string posing as document text.
    """
    try:
        from PIL import Image

        with Image.open(io.BytesIO(data)) as probe:
            probe.verify()
    except ImportError as exc:  # pragma: no cover - Pillow is a core dependency
        raise ParserUnavailableError("Image ingestion requires Pillow") from exc
    except Exception as exc:
        raise DocumentParseError(f"{filename}: not a readable image ({exc})") from exc

    has_vision = is_vision_provider(vision_provider)
    if not has_vision and not tesseract_available():
        raise OcrUnavailableError(
            f"{filename}: image ingestion needs OCR, but no OCR engine is available: "
            "install Tesseract (the tesseract binary plus the 'ocr' extra) or configure "
            "a vision-capable model provider"
        )
    if ocr_engine is None:
        from app.ocr.engine import OcrEngine

        ocr_engine = OcrEngine()
    result = await ocr_engine.extract(
        image_bytes=data,
        provider=vision_provider if has_vision else None,
        extract_fields=False,
        vision_fallback=has_vision,
    )
    text = (getattr(result, "raw_text", "") or "").strip()
    if not text:
        raise DocumentParseError(f"{filename}: no text could be extracted from the image")
    return text, str(getattr(result, "engine_used", "") or "")


def decode_text(data: bytes) -> str:
    """Decode an uploaded text file: UTF-8 (with or without BOM), UTF-16 with a
    BOM, else Windows-1252 — the common encodings of real-world exports."""
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def extract_upload_text(
    data: bytes, *, ext: str, filename: str, report: dict[str, Any] | None = None
) -> str:
    """The text of an uploaded file, by extension, or an exception — never garbage.

    Raises UnsupportedDocumentError (415), DocumentParseError (422: unreadable,
    empty or textless) or ParserUnavailableError (503: parser not installed).
    """
    ext = ext.lower().lstrip(".")
    if not data.strip():
        raise DocumentParseError(f"{filename}: the file is empty")
    if ext == "pdf":
        return "\n".join(extract_pdf_pages(data, filename=filename))
    if ext == "docx":
        return extract_docx_text(data, filename=filename)
    if ext == "pptx":
        return extract_pptx_text(data, filename=filename)
    if ext == "ppt":
        raise UnsupportedDocumentError(
            f"{filename}: legacy binary .ppt files are not supported; "
            "save the presentation as .pptx (or PDF) and upload that"
        )
    if ext in IMAGE_UPLOAD_EXTS:
        raise UnsupportedDocumentError(
            f"{filename}: images are OCR'd by the upload endpoint, not text extraction"
        )
    if ext in _UNSUPPORTED_BINARY_EXTS:
        raise UnsupportedDocumentError(
            f"{filename}: .{ext} files are not supported; "
            "convert to PDF, DOCX, XLSX, PPTX or text"
        )
    if ext in {"xlsx", "xlsm"}:
        from app.ingestion.parsers.excel_parser import ExcelParser

        text, excel_report = ExcelParser().parse_with_report(data, filename=filename)
        if report is not None:
            report.update(excel_report)  # a truncated workbook is reported to the caller
        if not text.strip() or all(
            line.startswith("Sheet:") for line in text.splitlines() if line.strip()
        ):
            raise DocumentParseError(f"{filename}: not a readable workbook, or it has no data")
        return text
    if b"\x00" in data[:8192] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise UnsupportedDocumentError(
            f"{filename}: unsupported binary file; upload PDF, DOCX, XLSX or text"
        )
    raw = decode_text(data)
    text = _extract_structured_text(raw, ext=ext, filename=filename)
    if not text.strip():
        raise DocumentParseError(f"{filename}: no extractable text")
    return text


def _extract_structured_text(raw: str, *, ext: str, filename: str) -> str:
    if ext in {"html", "htm", "xhtml"}:
        from app.ingestion.parsers.html_parser import HTMLParser

        return HTMLParser().parse(raw)
    if ext in {"csv", "tsv"}:
        from app.ingestion.parsers.csv_parser import CSVParser

        return CSVParser().parse(raw, filename=filename, delimiter="\t" if ext == "tsv" else "")
    if ext in {"json", "jsonl", "ndjson"}:
        from app.ingestion.parsers.json_parser import JSONParser

        return JSONParser().parse(raw, is_jsonl=ext != "json")
    if ext in {"yaml", "yml"}:
        from app.ingestion.parsers.yaml_parser import YAMLParser

        return YAMLParser().parse(raw, filename=filename)
    if ext == "ipynb":
        from app.ingestion.parsers.notebook_parser import NotebookParser

        return NotebookParser().parse(raw, filename=filename)
    if ext == "eml":
        from app.ingestion.parsers.email_parser import EmailParser

        return "\n\n".join(EmailParser().parse(raw))
    return raw
